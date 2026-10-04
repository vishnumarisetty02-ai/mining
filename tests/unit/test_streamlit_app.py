"""Focused tests for the SQLite-backed Streamlit reporting MVP."""

from __future__ import annotations

import json
import sqlite3
import uuid
import zipfile
from io import BytesIO
from pathlib import Path

import app as reporting
import pytest
from docx import Document as DocxDocument

EXPECTED_FACT_COUNT = 16
FACTS_PER_COMPARISON = 2


@pytest.fixture
def indexed_demo(monkeypatch: pytest.MonkeyPatch):
    db_path = reporting.DATA_DIR / f"test-{uuid.uuid4().hex}.sqlite3"
    monkeypatch.setattr(reporting, "DB_PATH", db_path)
    reporting.init_db()
    demo = Path(reporting.APP_DIR, "demo_data", "synthetic_mine_report.csv")
    inserted, fact_count = reporting.ingest_document(demo.read_bytes(), demo.name)
    assert inserted
    yield db_path, fact_count
    for path in (db_path, Path(f"{db_path}-wal"), Path(f"{db_path}-shm")):
        path.unlink(missing_ok=True)


def test_demo_ingestion_extracts_normalized_facts_and_metadata(indexed_demo):
    _, fact_count = indexed_demo
    facts = reporting.facts_df()
    documents = reporting.documents_df()

    assert fact_count == EXPECTED_FACT_COUNT
    assert len(facts) == EXPECTED_FACT_COUNT
    assert set(facts["period"]) == {"FY2023-24", "FY2024-25"}
    assert {
        "production",
        "target",
        "dispatch",
        "reserve",
        "gcv",
        "overburden",
        "stripping_ratio",
        "depth",
    } <= set(facts["metric"])
    production = facts[(facts["metric"] == "production") & (facts["period"] == "FY2024-25")].iloc[0]
    assert production["value"] == pytest.approx(1_350_000)
    assert production["original_value"] == pytest.approx(1.35)
    assert production["original_unit"] == "MT"
    assert production["mine"] == "Mine Alpha"
    assert production["subsidiary"] == "Synthetic Subsidiary"
    assert production["project"] == "Alpha Exploration"
    assert production["source_snippet"]
    assert production["fact_id"].startswith("FCT-")
    metadata = json.loads(documents.iloc[0]["metadata"])
    assert metadata["table_count"] == 1
    assert metadata["fact_count"] == EXPECTED_FACT_COUNT


def test_pdf_docx_xlsx_and_txt_extraction_are_supported():
    frame = reporting.pd.DataFrame(
        [
            {
                "Mine": "Mine Beta",
                "Reporting Period": "FY2024-25",
                "Production (MT)": 1.7,
            }
        ]
    )
    workbook = BytesIO()
    with reporting.pd.ExcelWriter(workbook, engine="openpyxl") as writer:
        frame.to_excel(writer, index=False)

    docx_document = DocxDocument()
    docx_document.add_paragraph("Synthetic geological report")
    table = docx_document.add_table(rows=2, cols=3)
    for column, heading in enumerate(frame.columns):
        table.cell(0, column).text = str(heading)
        table.cell(1, column).text = str(frame.iloc[0, column])
    docx_bytes = BytesIO()
    docx_document.save(docx_bytes)

    pdf_document = reporting.fitz.open()
    pdf_page = pdf_document.new_page()
    pdf_page.insert_text(
        (72, 72),
        "Mine: Mine Beta\nReporting Period: FY2024-25\nProduction: 1.7 MT",
    )
    pdf_bytes = pdf_document.tobytes()
    pdf_document.close()

    examples = [
        ("sample.pdf", pdf_bytes),
        ("sample.docx", docx_bytes.getvalue()),
        ("sample.xlsx", workbook.getvalue()),
        (
            "sample.txt",
            b"Mine: Mine Beta\nReporting Period: FY2024-25\nProduction: 1.7 MT",
        ),
    ]
    for filename, content in examples:
        pages, extracted_text, tables = reporting.extract_document(content, filename)
        facts = reporting.extract_facts(pages, reporting.sha256(content))
        assert extracted_text
        assert tables or filename in {"sample.pdf", "sample.txt"}
        assert any(
            fact["metric"] == "production"
            and fact["entity"] == "Mine Beta"
            and fact["period"] == "FY2024-25"
            for fact in facts
        ), filename


def test_domain_ontology_aliases_drive_fact_extraction():
    reporting.ONTOLOGY["roof_thickness"] = {
        "label": "Roof Thickness",
        "aliases": ["roof thickness"],
        "units": ["cm"],
    }
    try:
        text = (
            "Mine: Mine Gamma\nReporting Period: FY2024-25\n"
            "Seam thickness: 3.1 m\nAsh content: 8 percent\n"
            "Roof thickness: 160 cm"
        )
        facts = reporting.extract_facts(
            [{"page": 1, "text": text}], reporting.sha256(text.encode())
        )
        extracted = {fact["metric"]: fact for fact in facts}
        assert extracted["seam_thickness"]["value"] == pytest.approx(3.1)
        assert extracted["ash"]["unit"] == "percent"
        assert extracted["roof_thickness"]["unit"] == "cm"
        assert extracted["roof_thickness"]["value"] == pytest.approx(160)
    finally:
        reporting.ONTOLOGY.pop("roof_thickness", None)


def test_query_change_and_parliamentary_draft_use_fact_evidence(indexed_demo):
    assert indexed_demo[1] == EXPECTED_FACT_COUNT
    result = reporting.evidence_answer("What was Mine Alpha production in FY2024-25?")
    assert result["fact"]["value"] == pytest.approx(1_350_000)
    assert result["fact"]["fact_id"] in result["answer"]
    assert "1,350,000 tonnes" in result["answer"]

    comparison_answer = reporting.evidence_answer(
        "Compare Mine Alpha production between FY2023-24 and FY2024-25"
    )
    assert len(comparison_answer["facts"]) == FACTS_PER_COMPARISON
    assert "12.50%" in comparison_answer["answer"]

    reserve_change = reporting.evidence_answer(
        "Show reserve changes between FY2023-24 and FY2024-25"
    )
    assert len(reserve_change["comparisons"]) == 1
    assert reserve_change["comparisons"][0]["absolute_change"] == pytest.approx(600_000)

    highest = reporting.evidence_answer("Which mine has the highest production?")
    assert highest["fact"]["entity"] == "Mine Alpha"
    assert highest["fact"]["period"] == "FY2024-25"

    unanswered = reporting.evidence_answer("What was Mine Unknown production in FY2024-25?")
    assert unanswered["fact"] is None
    assert "No numerical value" in unanswered["answer"]

    analysis, draft, answer = reporting.draft_parliamentary_answer(
        "State Mine Alpha production in FY2023-24 and FY2024-25 and the percentage change."
    )
    assert analysis["metrics"] == ["production"]
    assert analysis["entities"] == ["Mine Alpha"]
    assert len(answer["facts"]) == FACTS_PER_COMPARISON
    assert "DRAFT" in draft
    assert all(fact["fact_id"] in draft for fact in answer["facts"])


def test_change_intelligence_returns_delta_and_field_status(indexed_demo):
    assert indexed_demo[1] == EXPECTED_FACT_COUNT
    changes = reporting.compare_entity_periods("Mine Alpha", "FY2023-24", "FY2024-25")
    production = changes[changes["Metric"] == "production"].iloc[0]
    assert production["Absolute Change"] == pytest.approx(150_000)
    assert production["% Change"] == pytest.approx(12.5)
    assert production["Status"] == "CHANGED"
    assert production["Attention"] == "NORMAL"


def test_report_outputs_and_authorized_approval_are_audited(indexed_demo, monkeypatch):
    assert indexed_demo[1] == EXPECTED_FACT_COUNT
    selected = reporting.facts_df()
    firewall = reporting.consistency_firewall()
    assert firewall["status"] == "PASS"

    docx_data = reporting.generate_docx("Synthetic test report", selected, firewall)
    with zipfile.ZipFile(BytesIO(docx_data)) as package:
        assert "word/document.xml" in package.namelist()

    pdf_data = reporting.generate_pdf("Synthetic test report", selected, firewall)
    assert pdf_data.startswith(b"%PDF")

    report_id = reporting.record_generated_report(
        "Synthetic test report", "PDF", firewall, selected
    )
    monkeypatch.setenv("CMPDI_APPROVER_PIN", "test-secret")
    assert not reporting.verify_fact(selected.iloc[0]["fact_id"], "Test Approver", "incorrect")
    assert reporting.verify_fact(selected.iloc[0]["fact_id"], "Test Approver", "test-secret")
    assert not reporting.approve_report(report_id, "Test Approver", "incorrect")
    assert reporting.approve_report(report_id, "Test Approver", "test-secret")

    conn = sqlite3.connect(reporting.DB_PATH)
    conn.row_factory = sqlite3.Row
    try:
        approved_row = conn.execute("SELECT * FROM reports WHERE id=?", (report_id,)).fetchone()
        event_count = conn.execute(
            "SELECT COUNT(*) FROM audits WHERE event_type='REPORT_APPROVED'"
        ).fetchone()[0]
    finally:
        conn.close()
    assert (approved_row["status"], approved_row["approved_by"]) == ("Approved", "Test Approver")
    assert event_count == 1
    final_note = f"APPROVED by {approved_row['approved_by']} on {approved_row['approved_at']}."
    final_docx = reporting.generate_docx(
        approved_row["title"],
        selected,
        json.loads(approved_row["firewall_data"]),
        approval_note=final_note,
    )
    with zipfile.ZipFile(BytesIO(final_docx)) as package:
        document_text = package.read("word/document.xml").decode()
    assert "APPROVED by Test Approver" in document_text


def test_firewall_blocks_conflicting_source_values(indexed_demo, monkeypatch):
    assert indexed_demo[1] == EXPECTED_FACT_COUNT
    monkeypatch.setenv("CMPDI_APPROVER_PIN", "test-secret")
    facts = reporting.facts_df()
    production = facts[(facts["metric"] == "production") & (facts["period"] == "FY2024-25")].iloc[0]
    conflicting_fact = {
        "fact_id": "FCT-CONFLICT-TEST",
        "entity": production["entity"],
        "mine": production["mine"],
        "subsidiary": production["subsidiary"],
        "project": production["project"],
        "metric": production["metric"],
        "period": production["period"],
        "value": float(production["value"]) + 50_000,
        "unit": production["unit"],
        "original_value": 1.4,
        "original_unit": "MT",
        "source_page": 1,
        "source_snippet": "Synthetic conflicting production value.",
        "extraction_method": "Test",
        "confidence": 0.95,
    }
    reporting.save_document(
        b"synthetic-conflict-source",
        "synthetic_conflict.txt",
        "txt",
        "Synthetic conflicting production value.",
        [conflicting_fact],
    )
    firewall = reporting.consistency_firewall()

    assert firewall["status"] == "REVIEW_REQUIRED"
    conflict = next(item for item in firewall["findings"] if item["type"] == "SOURCE_CONFLICT")
    assert "synthetic_conflict.txt" in conflict["detail"]
    report_id = reporting.record_generated_report("Conflicted draft", "DOCX", firewall, facts)
    assert not reporting.approve_report(report_id, "Test Approver", "test-secret")


def test_firewall_recalculates_reported_achievement(indexed_demo):  # noqa: ARG001
    source = (
        b"Mine: Mine Beta\nReporting Period: FY2024-25\n"
        b"Production: 80 tonnes\nTarget: 100 tonnes\nAchievement: 70%"
    )
    pages, text, tables = reporting.extract_document(source, "achievement_check.txt")
    facts = reporting.extract_facts(pages, reporting.sha256(source))
    reporting.save_document(source, "achievement_check.txt", "txt", text, facts, tables)

    firewall = reporting.consistency_firewall()
    finding = next(
        item for item in firewall["findings"] if item["type"] == "ARITHMETIC_INCONSISTENCY"
    )
    assert firewall["status"] == "REVIEW_REQUIRED"
    assert "70.00%" in finding["message"]
    assert "80.00%" in finding["message"]
