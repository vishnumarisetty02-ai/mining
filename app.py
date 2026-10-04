# ============================================================
# CMPDI / CIL UNIQUE AI REPORTING INTELLIGENCE
# ============================================================
# Smart Automation MVP
#
# NEW / DIFFERENTIATING FEATURES
# -------------------------------
# 1. Fact Passport - every extracted fact gets provenance
# 2. Consistency Firewall - blocks suspicious/inconsistent reports
# 3. Change Intelligence - compare reporting periods and detect changes
# 4. Parliamentary Question Copilot - turns a question into a sourced draft
# 5. Mining Domain Ontology - maps terms to domain entities/metrics
# 6. Evidence Trace - answer -> fact -> source document -> page
# 7. Topic Evolution - topic frequencies by reporting period
# 8. Human approval workflow
# 9. Automated DOCX/PDF report generation
# 10. Optional local Ollama wording model
#
# Run:
#   pip install -r requirements.txt
#   streamlit run app.py
# ============================================================

from __future__ import annotations

import hashlib
import html
import io
import json
import math
import os
import re
import secrets
import sqlite3
from datetime import datetime
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
import streamlit as st

# ---------------- Optional libraries ----------------
try:
    import pymupdf as fitz
except Exception:
    fitz = None

try:
    from docx import Document
except Exception:
    Document = None

try:
    from PIL import Image
except Exception:
    Image = None

try:
    import pytesseract
except Exception:
    pytesseract = None

try:
    from sklearn.feature_extraction.text import TfidfVectorizer
    from sklearn.metrics.pairwise import cosine_similarity
except Exception:
    TfidfVectorizer = None
    cosine_similarity = None

APP_DIR = Path(__file__).resolve().parent
DATA_DIR = APP_DIR / "data"
DATA_DIR.mkdir(exist_ok=True)
os.environ.setdefault("MPLCONFIGDIR", str(APP_DIR / ".matplotlib-cache"))

try:
    from wordcloud import WordCloud
except Exception:
    WordCloud = None

try:
    import matplotlib.pyplot as plt
except Exception:
    plt = None

try:
    import requests
except Exception:
    requests = None

try:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
except Exception:
    SimpleDocTemplate = None


DB_PATH = Path(os.environ.get("CMPDI_DB_PATH", str(DATA_DIR / "cmpdi_unique_reporting.db")))
ONTOLOGY_PATH = APP_DIR / "models" / "mining_ontology.json"


# ============================================================
# CONFIG / DOMAIN ONTOLOGY
# ============================================================

ONTOLOGY = {
    "production": {
        "label": "Coal Production",
        "aliases": ["production", "coal production", "output", "produced", "coal output"],
        "units": ["tonnes", "MT", "million tonnes", "lakh tonnes"],
    },
    "target": {
        "label": "Production Target",
        "aliases": ["target", "production target"],
        "units": ["tonnes", "MT", "million tonnes", "lakh tonnes"],
    },
    "dispatch": {
        "label": "Coal Dispatch",
        "aliases": ["dispatch", "despatch", "coal dispatch", "dispatches"],
        "units": ["tonnes", "MT", "million tonnes", "lakh tonnes"],
    },
    "reserve": {
        "label": "Coal Reserve",
        "aliases": ["reserve", "reserves", "geological reserve", "resource"],
        "units": ["tonnes", "MT", "million tonnes", "lakh tonnes"],
    },
    "gcv": {
        "label": "Gross Calorific Value",
        "aliases": ["gcv", "gross calorific value", "calorific value"],
        "units": ["kcal/kg"],
    },
    "overburden": {
        "label": "Overburden Removal",
        "aliases": ["overburden", "ob removal", "ob", "over burden"],
        "units": ["MCum", "million m3", "m3"],
    },
    "stripping_ratio": {
        "label": "Stripping Ratio",
        "aliases": ["stripping ratio", "sr"],
        "units": ["ratio"],
    },
    "seam_thickness": {
        "label": "Coal Seam Thickness",
        "aliases": ["seam thickness", "coal seam thickness"],
        "units": ["m", "metres"],
    },
    "ash": {
        "label": "Ash Content",
        "aliases": ["ash content", "ash percentage"],
        "units": ["%", "percent"],
    },
    "moisture": {
        "label": "Moisture Content",
        "aliases": ["moisture", "moisture content"],
        "units": ["%", "percent"],
    },
    "depth": {
        "label": "Depth",
        "aliases": ["depth", "drilling depth", "borehole depth"],
        "units": ["m", "metre", "metres"],
    },
}

UNIT_FACTORS_TO_TONNES = {
    "tonne": 1.0,
    "tonnes": 1.0,
    "t": 1.0,
    "mt": 1_000_000.0,
    "million tonnes": 1_000_000.0,
    "lakh tonnes": 100_000.0,
}


STOPWORDS = {
    "the",
    "and",
    "for",
    "with",
    "from",
    "this",
    "that",
    "were",
    "have",
    "has",
    "had",
    "are",
    "was",
    "will",
    "into",
    "during",
    "their",
    "there",
    "which",
    "about",
    "been",
    "also",
    "than",
    "such",
    "where",
    "what",
    "report",
    "page",
    "table",
    "mine",
    "project",
    "year",
    "years",
    "shall",
    "would",
    "could",
    "should",
    "based",
    "following",
}


# ============================================================
# DB
# ============================================================


def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def init_db():
    conn = get_db()
    cur = conn.cursor()

    cur.execute("""
        CREATE TABLE IF NOT EXISTS documents(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            sha256 TEXT UNIQUE NOT NULL,
            filename TEXT NOT NULL,
            file_type TEXT NOT NULL,
            uploaded_at TEXT NOT NULL,
            content TEXT NOT NULL,
            metadata TEXT DEFAULT '{}'
        )
    """)

    cur.execute("""
        CREATE TABLE IF NOT EXISTS facts(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            fact_id TEXT UNIQUE NOT NULL,
            document_id INTEGER NOT NULL,
            entity TEXT NOT NULL,
            mine TEXT NOT NULL DEFAULT 'Unknown',
            subsidiary TEXT NOT NULL DEFAULT 'Unknown',
            project TEXT NOT NULL DEFAULT 'Unknown',
            metric TEXT NOT NULL,
            period TEXT NOT NULL,
            value REAL NOT NULL,
            unit TEXT NOT NULL,
            original_value REAL,
            original_unit TEXT,
            source_page INTEGER,
            source_snippet TEXT,
            extraction_method TEXT,
            confidence REAL NOT NULL,
            verification_status TEXT NOT NULL DEFAULT 'Pending',
            created_at TEXT NOT NULL,
            FOREIGN KEY(document_id) REFERENCES documents(id)
        )
    """)

    cur.execute("""
        CREATE TABLE IF NOT EXISTS audits(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            event_type TEXT NOT NULL,
            message TEXT NOT NULL,
            created_at TEXT NOT NULL
        )
    """)

    cur.execute("""
        CREATE TABLE IF NOT EXISTS reports(
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            title TEXT NOT NULL,
            file_type TEXT NOT NULL,
            status TEXT NOT NULL DEFAULT 'Draft',
            generated_at TEXT NOT NULL,
            approved_at TEXT,
            approved_by TEXT,
            firewall_status TEXT NOT NULL,
            fact_ids TEXT NOT NULL,
            firewall_data TEXT NOT NULL DEFAULT '{}',
            comparison_data TEXT NOT NULL DEFAULT '[]'
        )
    """)

    existing_columns = {row["name"] for row in cur.execute("PRAGMA table_info(facts)").fetchall()}
    for column in ("mine", "subsidiary", "project"):
        if column not in existing_columns:
            cur.execute(f"ALTER TABLE facts ADD COLUMN {column} TEXT NOT NULL DEFAULT 'Unknown'")
    report_columns = {row["name"] for row in cur.execute("PRAGMA table_info(reports)").fetchall()}
    for column, default in (
        ("firewall_data", "'{}'"),
        ("comparison_data", "'[]'"),
    ):
        if column not in report_columns:
            cur.execute(f"ALTER TABLE reports ADD COLUMN {column} TEXT NOT NULL DEFAULT {default}")

    conn.commit()
    conn.close()


def audit(event_type: str, message: str):
    conn = get_db()
    conn.execute(
        "INSERT INTO audits(event_type,message,created_at) VALUES(?,?,?)",
        (event_type, message, datetime.now().isoformat(timespec="seconds")),
    )
    conn.commit()
    conn.close()


# ============================================================
# HELPERS
# ============================================================


def sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def clean_number(s: str) -> float:
    return float(s.replace(",", "").strip())


def normalize_period(text: str) -> str:
    s = str(text).strip()
    m = re.search(r"FY\s*(20\d{2})\s*[-–/]\s*(\d{2}|20\d{2})", s, re.I)
    if m:
        y1 = int(m.group(1))
        y2 = m.group(2)
        return f"FY{y1}-{y2[-2:]}"
    m = re.search(r"(20\d{2})\s*[-–/]\s*(\d{2}|20\d{2})", s)
    if m:
        return f"FY{m.group(1)}-{m.group(2)[-2:]}"
    return "Unknown Period"


def detect_period(text: str) -> str:
    for pattern in [
        r"FY\s*20\d{2}\s*[-–/]\s*(?:\d{2}|20\d{2})",
        r"20\d{2}\s*[-–/]\s*(?:\d{2}|20\d{2})",
    ]:
        m = re.search(pattern, text, re.I)
        if m:
            return normalize_period(m.group(0))
    return "Unknown Period"


def detect_entity(text: str) -> str:
    patterns = [
        r"(?im)^\s*(?:Mine|Project|Block|Unit)\s*[:\-]\s*([^\r\n,;|]{2,70})",
        r"(?im)^\s*Subsidiary\s*[:\-]\s*([^\r\n,;|]{2,50})",
        r"\b([A-Z][A-Za-z0-9 &.-]{2,50})\s+(?:OC|OCP|UG|Opencast|Open Cast|Underground)\b",
        r"\bMine\s+([A-Z][A-Za-z0-9-]*(?:\s+[A-Z][A-Za-z0-9-]*){0,3})",
    ]
    for pat in patterns:
        m = re.search(pat, text)
        if m:
            return m.group(1).strip().rstrip(".,;:-")
    return "Unknown Entity"


def detect_named_field(text: str, field: str) -> str:
    match = re.search(
        rf"(?im)^\s*{re.escape(field)}\s*[:\-]\s*([^\r\n,;|]{{2,70}})",
        text,
    )
    return match.group(1).strip().rstrip(".,;:-") if match else "Unknown"


def normalize_tonnes(value: float, unit: str) -> tuple[float, str]:
    u = unit.lower().strip()
    if u in UNIT_FACTORS_TO_TONNES:
        return value * UNIT_FACTORS_TO_TONNES[u], "tonnes"
    return value, unit


def fact_id_for(
    doc_hash: str, metric: str, period: str, value: float, page: Any, entity: str = ""
) -> str:
    raw = f"{doc_hash}|{metric}|{period}|{value:.8f}|{page}|{entity}"
    return "FCT-" + hashlib.sha256(raw.encode()).hexdigest()[:12].upper()


def confidence_for(snippet: str, metric: str, unit: str, page: Any) -> float:
    score = 0.78
    s = snippet.lower()
    if metric in s:
        score += 0.07
    if unit:
        score += 0.05
    if page is not None:
        score += 0.04
    if any(k in s for k in ["table", "statement", "summary"]):
        score += 0.04
    return min(score, 0.98)


def render_table_rows(table: pd.DataFrame) -> str:
    known_units = sorted(
        {str(unit) for item in ONTOLOGY.values() for unit in item["units"]},
        key=len,
        reverse=True,
    )
    unit_expression = "|".join(re.escape(unit) for unit in known_units)
    rows = []
    for _, row in table.iterrows():
        cells = []
        for column, value in row.items():
            if pd.isna(value):
                continue
            heading = str(column).strip()
            is_ontology_term = any(
                heading.lower()
                in {
                    metric.replace("_", " ").lower(),
                    item["label"].lower(),
                    *(alias.lower() for alias in item["aliases"]),
                }
                for metric, item in ONTOLOGY.items()
            )
            unit_match = (
                re.search(rf"(?<!\w)(?:{unit_expression})(?!\w)", heading, re.I)
                if unit_expression and not is_ontology_term
                else None
            )
            clean_heading = re.sub(
                rf"\s*\(\s*(?:{unit_expression})\s*\)",
                "",
                heading,
                flags=re.I,
            ).strip()
            if not is_ontology_term:
                clean_heading = re.sub(
                    rf"\s+(?:{unit_expression})(?=$|[^\w])",
                    "",
                    clean_heading,
                    flags=re.I,
                ).strip()
            suffix = f" {unit_match.group(0)}" if unit_match else ""
            cells.append(f"{clean_heading}: {value}{suffix}")
        if cells:
            rows.append("\n".join(cells))
    return "\n\n".join(rows)


# ============================================================
# EXTRACTION
# ============================================================


def extract_pdf(data: bytes, filename: str):
    if fitz is None:
        raise RuntimeError("Install PyMuPDF first.")

    doc = fitz.open(stream=data, filetype="pdf")
    pages = []
    all_text = []

    for page_num, page in enumerate(doc, start=1):
        txt = page.get_text("text").strip()

        method = "Text/Parser"
        if not txt and pytesseract is not None and Image is not None:
            pix = page.get_pixmap(matrix=fitz.Matrix(1.6, 1.6), alpha=False)
            image = Image.open(io.BytesIO(pix.tobytes("png")))
            try:
                txt = pytesseract.image_to_string(image)
                method = "OCR"
            except Exception as exc:
                raise RuntimeError(
                    f"OCR failed on page {page_num}; verify that Tesseract OCR is installed."
                ) from exc
        elif not txt:
            raise RuntimeError(
                f"No selectable text on PDF page {page_num}. Install Tesseract OCR "
                "and configure pytesseract to process scanned pages."
            )

        pages.append({"page": page_num, "text": txt, "extraction_method": method})
        all_text.append(txt)

    doc.close()
    return pages, "\n\n".join(all_text)


def extract_docx(data: bytes):
    if Document is None:
        raise RuntimeError("Install python-docx first.")
    doc = Document(io.BytesIO(data))
    text_parts = [p.text for p in doc.paragraphs if p.text.strip()]
    tables = []

    for table in doc.tables:
        rows = []
        for row in table.rows:
            rows.append([c.text.strip() for c in row.cells])
        if rows:
            frame = pd.DataFrame(rows[1:], columns=rows[0]) if len(rows) > 1 else pd.DataFrame(rows)
            tables.append(frame)

    table_text = "\n\n".join(render_table_rows(table) for table in tables)
    txt = "\n".join(text_parts + ([table_text] if table_text else []))
    return [{"page": None, "text": txt, "extraction_method": "DOCX Parser"}], txt, tables


def extract_xlsx(data: bytes):
    xls = pd.ExcelFile(io.BytesIO(data))
    blocks = []
    tables = []
    for sheet in xls.sheet_names:
        df = pd.read_excel(xls, sheet_name=sheet)
        tables.append(df)
        blocks.append(f"Sheet: {sheet}\n{render_table_rows(df)}")
    txt = "\n\n".join(blocks)
    return [{"page": None, "text": txt, "extraction_method": "XLSX Table Parser"}], txt, tables


def extract_csv(data: bytes):
    df = pd.read_csv(io.BytesIO(data))
    txt = render_table_rows(df)
    return [{"page": None, "text": txt, "extraction_method": "CSV Table Parser"}], txt, [df]


def extract_text(data: bytes):
    txt = data.decode("utf-8", errors="ignore")
    return [{"page": None, "text": txt, "extraction_method": "Text Parser"}], txt, []


def extract_image(data: bytes):
    if Image is None or pytesseract is None:
        raise RuntimeError("Install Pillow and pytesseract for image OCR.")
    image = Image.open(io.BytesIO(data))
    try:
        txt = pytesseract.image_to_string(image)
    except Exception as exc:
        raise RuntimeError(
            "Image OCR failed. Verify that Tesseract OCR is installed and configured."
        ) from exc
    if not txt.strip():
        raise RuntimeError("Image OCR completed but found no readable text.")
    return [{"page": 1, "text": txt, "extraction_method": "OCR"}], txt, []


def extract_document(data: bytes, filename: str):
    ext = Path(filename).suffix.lower()

    if ext == ".pdf":
        pages, text = extract_pdf(data, filename)
        return pages, text, []

    if ext == ".docx":
        return extract_docx(data)

    if ext == ".xlsx":
        return extract_xlsx(data)

    if ext == ".csv":
        return extract_csv(data)

    if ext == ".txt":
        return extract_text(data)

    if ext in [".png", ".jpg", ".jpeg", ".webp"]:
        return extract_image(data)

    raise ValueError(f"Unsupported file type: {ext}")


def extract_facts(pages: list[dict[str, Any]], doc_hash: str):
    facts = []

    patterns = {
        "production": [
            r"production(?:\s+of\s+coal)?\s*(?:is|was|=|:)?\s*([\d,]+(?:\.\d+)?)\s*(MT|million tonnes|lakh tonnes|tonnes|t)\b",
        ],
        "target": [
            r"target\s*(?:is|was|=|:)?\s*([\d,]+(?:\.\d+)?)\s*(MT|million tonnes|lakh tonnes|tonnes|t)\b",
        ],
        "dispatch": [
            r"(?:dispatch|despatch)\s*(?:is|was|=|:)?\s*([\d,]+(?:\.\d+)?)\s*(MT|million tonnes|lakh tonnes|tonnes|t)\b",
        ],
        "reserve": [
            r"(?:reserve|reserves|geological reserve)\s*(?:is|was|=|:)?\s*([\d,]+(?:\.\d+)?)\s*(MT|million tonnes|lakh tonnes|tonnes|t)\b",
        ],
        "gcv": [
            r"(?:GCV|gross calorific value)\s*(?:is|was|=|:)?\s*([\d,]+(?:\.\d+)?)\s*(kcal/kg)\b",
        ],
        "overburden": [
            r"(?:overburden|OB removal)\s*(?:is|was|=|:)?\s*([\d,]+(?:\.\d+)?)\s*(MCum|million m3|m3)\b",
        ],
        "depth": [
            r"(?:depth|drilling depth|borehole depth)\s*(?:is|was|=|:)?\s*([\d,]+(?:\.\d+)?)\s*(m|metre|metres)\b",
        ],
        "stripping_ratio": [
            r"(?:stripping\s+ratio|\bSR\b)\s*(?:is|was|=|:)?\s*([\d,]+(?:\.\d+)?)\s*(?:ratio|:1)?\b",
        ],
    }
    for metric, meta in ONTOLOGY.items():
        aliases = sorted(
            {metric.replace("_", " "), meta["label"], *meta["aliases"]},
            key=len,
            reverse=True,
        )
        units = sorted(set(meta["units"]), key=len, reverse=True)
        if not aliases or not units:
            continue
        alias_expression = "|".join(re.escape(alias) for alias in aliases)
        unit_expression = "|".join(re.escape(unit) for unit in units)
        patterns.setdefault(metric, []).append(
            rf"(?<!\w)(?:{alias_expression})(?!\w)"
            rf"\s*(?:is|was|=|:)?\s*([\d,]+(?:\.\d+)?)"
            rf"\s*({unit_expression})(?=$|[^\w])"
        )

    for pg in pages:
        page_num = pg.get("page")
        text = pg.get("text", "")
        if not text:
            continue

        page_period = detect_period(text)
        entity = detect_entity(text)
        mine = detect_named_field(text, "Mine")
        subsidiary = detect_named_field(text, "Subsidiary")
        project = detect_named_field(text, "Project")
        if entity == "Unknown Entity":
            entity = (
                mine if mine != "Unknown" else (project if project != "Unknown" else subsidiary)
            )

        for metric, plist in patterns.items():
            for pat in plist:
                for match in re.finditer(pat, text, flags=re.I):
                    value_raw = match.group(1)
                    unit_raw = (
                        match.group(2)
                        if match.lastindex is not None and match.lastindex >= 2
                        else ""
                    )
                    if metric == "stripping_ratio" and not unit_raw:
                        unit_raw = "ratio"
                    try:
                        original_value = clean_number(value_raw)
                    except ValueError:
                        continue

                    value = original_value
                    unit = unit_raw

                    if metric in {"production", "target", "dispatch", "reserve"}:
                        value, unit = normalize_tonnes(value, unit)

                    record_start = text.rfind("\n\n", 0, match.start()) + 2
                    record_end = text.find("\n\n", match.end())
                    if record_end < 0:
                        record_end = len(text)
                    record_text = text[record_start:record_end]
                    record_period = detect_period(record_text)
                    fact_period = (
                        record_period if record_period != "Unknown Period" else page_period
                    )
                    fact_mine = detect_named_field(record_text, "Mine")
                    fact_subsidiary = detect_named_field(record_text, "Subsidiary")
                    fact_project = detect_named_field(record_text, "Project")
                    if fact_mine == "Unknown":
                        fact_mine = mine
                    if fact_subsidiary == "Unknown":
                        fact_subsidiary = subsidiary
                    if fact_project == "Unknown":
                        fact_project = project
                    fact_entity = detect_entity(record_text)
                    if fact_mine != "Unknown":
                        fact_entity = fact_mine
                    elif fact_entity == "Unknown Entity":
                        fact_entity = entity
                    if fact_entity == "Unknown Entity":
                        fact_entity = next(
                            (
                                value
                                for value in (fact_mine, fact_project, fact_subsidiary)
                                if value != "Unknown"
                            ),
                            fact_entity,
                        )

                    snippet = text[max(0, match.start() - 100) : min(len(text), match.end() + 120)]
                    conf = confidence_for(snippet, metric, unit, page_num)

                    fid = fact_id_for(doc_hash, metric, fact_period, value, page_num, fact_entity)

                    facts.append(
                        {
                            "fact_id": fid,
                            "entity": fact_entity,
                            "mine": fact_mine,
                            "subsidiary": fact_subsidiary,
                            "project": fact_project,
                            "metric": metric,
                            "period": fact_period,
                            "value": float(value),
                            "unit": unit,
                            "original_value": float(original_value),
                            "original_unit": unit_raw,
                            "source_page": page_num,
                            "source_snippet": " ".join(snippet.split()),
                            "extraction_method": pg.get("extraction_method", "Text/Parser"),
                            "confidence": round(conf, 3),
                        }
                    )

    # Remove exact duplicates
    unique = {}
    for f in facts:
        key = (
            f["entity"],
            f["metric"],
            f["period"],
            f["value"],
            f["unit"],
            f["source_page"],
            f["entity"],
        )
        unique[key] = f
    return list(unique.values())


def save_document(
    data: bytes,
    filename: str,
    file_type: str,
    text: str,
    facts: list[dict[str, Any]],
    tables: list[pd.DataFrame] | None = None,
):
    digest = sha256(data)
    conn = get_db()

    cur = conn.cursor()
    cur.execute("SELECT id FROM documents WHERE sha256=?", (digest,))
    old = cur.fetchone()
    if old:
        conn.close()
        return int(old["id"]), False

    cur.execute(
        """
        INSERT INTO documents(sha256,filename,file_type,uploaded_at,content,metadata)
        VALUES(?,?,?,?,?,?)
    """,
        (
            digest,
            filename,
            file_type,
            datetime.now().isoformat(timespec="seconds"),
            text,
            json.dumps(
                {
                    "fact_count": len(facts),
                    "table_count": len(tables or []),
                    "tables": [
                        table.head(100).to_dict(orient="records") for table in (tables or [])
                    ],
                },
                default=str,
            ),
        ),
    )
    document_id = cur.lastrowid

    for f in facts:
        cur.execute(
            """
            INSERT OR IGNORE INTO facts(
                fact_id,document_id,entity,mine,subsidiary,project,metric,period,value,unit,
                original_value,original_unit,source_page,source_snippet,
                extraction_method,confidence,verification_status,created_at
            ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
        """,
            (
                f["fact_id"],
                document_id,
                f["entity"],
                f.get("mine", "Unknown"),
                f.get("subsidiary", "Unknown"),
                f.get("project", "Unknown"),
                f["metric"],
                f["period"],
                f["value"],
                f["unit"],
                f["original_value"],
                f["original_unit"],
                f["source_page"],
                f["source_snippet"],
                f["extraction_method"],
                f["confidence"],
                "Pending",
                datetime.now().isoformat(timespec="seconds"),
            ),
        )

    conn.commit()
    conn.close()
    audit("DOCUMENT_INGESTED", f"{filename}: {len(facts)} facts")
    return int(document_id), True


def ingest_document(data: bytes, filename: str) -> tuple[bool, int]:
    pages, text, tables = extract_document(data, filename)
    if not text.strip():
        raise ValueError(f"{filename}: no readable text or table content was extracted.")
    facts = extract_facts(pages, sha256(data))
    audit("FACTS_EXTRACTED", f"{filename}: {len(facts)} structured facts")
    file_type = Path(filename).suffix.lower().replace(".", "") or "unknown"
    _, inserted = save_document(data, filename, file_type, text, facts, tables)
    return inserted, len(facts)


# ============================================================
# DATA ACCESS
# ============================================================


def documents_df():
    conn = get_db()
    df = pd.read_sql_query("SELECT * FROM documents ORDER BY id DESC", conn)
    conn.close()
    return df


def facts_df():
    conn = get_db()
    df = pd.read_sql_query(
        """
        SELECT f.*, d.filename
        FROM facts f
        JOIN documents d ON d.id = f.document_id
        ORDER BY f.id DESC
    """,
        conn,
    )
    conn.close()
    return df


def is_authorized_approver(name: str, pin: str) -> bool:
    configured_pin = os.environ.get("CMPDI_APPROVER_PIN", "")
    return bool(name.strip() and configured_pin and secrets.compare_digest(pin, configured_pin))


def verify_fact(fact_id: str, name: str, pin: str) -> bool:
    if not is_authorized_approver(name, pin):
        return False
    conn = get_db()
    cur = conn.execute(
        """UPDATE facts SET verification_status='Verified'
           WHERE fact_id=? AND verification_status!='Verified'""",
        (fact_id,),
    )
    conn.commit()
    updated = cur.rowcount == 1
    conn.close()
    if updated:
        audit("FACT_VERIFIED", f"{fact_id} verified by {name.strip()}")
    return updated


def record_generated_report(
    title: str,
    file_type: str,
    firewall: dict[str, Any],
    selected: pd.DataFrame,
    comparison: pd.DataFrame | None = None,
) -> int:
    fact_ids = selected["fact_id"].astype(str).tolist() if not selected.empty else []
    comparison_data = (
        comparison.to_json(orient="records", date_format="iso")
        if comparison is not None and not comparison.empty
        else "[]"
    )
    conn = get_db()
    cur = conn.execute(
        """INSERT INTO reports(
               title,file_type,status,generated_at,firewall_status,fact_ids,
               firewall_data,comparison_data
           ) VALUES(?,?,'Draft',?,?,?,?,?)""",
        (
            title,
            file_type,
            datetime.now().isoformat(timespec="seconds"),
            firewall["status"],
            json.dumps(fact_ids),
            json.dumps(firewall),
            comparison_data,
        ),
    )
    report_id = int(cur.lastrowid)
    conn.commit()
    conn.close()
    audit("REPORT_GENERATED", f"{file_type} draft #{report_id}: {title}")
    return report_id


def approve_report(report_id: int, name: str, pin: str) -> bool:
    if not is_authorized_approver(name, pin):
        return False
    conn = get_db()
    row = conn.execute(
        "SELECT firewall_status,status FROM reports WHERE id=?", (report_id,)
    ).fetchone()
    if row is None or row["firewall_status"] != "PASS" or row["status"] != "Draft":
        conn.close()
        return False
    conn.execute(
        """UPDATE reports SET status='Approved',approved_at=?,approved_by=?
           WHERE id=? AND status='Draft' AND firewall_status='PASS'""",
        (datetime.now().isoformat(timespec="seconds"), name.strip(), report_id),
    )
    conn.commit()
    updated = conn.total_changes > 0
    conn.close()
    if updated:
        audit("REPORT_APPROVED", f"Report #{report_id} approved by {name.strip()}")
    return updated


def save_ontology() -> None:
    ONTOLOGY_PATH.parent.mkdir(parents=True, exist_ok=True)
    ONTOLOGY_PATH.write_text(json.dumps(ONTOLOGY, indent=2), encoding="utf-8")


def load_ontology_overrides() -> None:
    if not ONTOLOGY_PATH.exists():
        return
    try:
        configured = json.loads(ONTOLOGY_PATH.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise RuntimeError(f"Could not load mining ontology: {exc}") from exc
    if not isinstance(configured, dict):
        raise ValueError("Mining ontology configuration must be a JSON object.")
    for key, entry in configured.items():
        if (
            isinstance(entry, dict)
            and isinstance(entry.get("label"), str)
            and isinstance(entry.get("aliases"), list)
            and isinstance(entry.get("units"), list)
        ):
            ONTOLOGY[key] = entry


# ============================================================
# UNIQUE FEATURE 1: FACT PASSPORT
# ============================================================


def build_fact_passport(row: pd.Series) -> dict[str, Any]:
    return {
        "FACT_ID": row["fact_id"],
        "Entity": row["entity"],
        "Mine": row.get("mine", "Unknown"),
        "Subsidiary": row.get("subsidiary", "Unknown"),
        "Project": row.get("project", "Unknown"),
        "Metric": row["metric"],
        "Period": row["period"],
        "Normalized Value": float(row["value"]),
        "Normalized Unit": row["unit"],
        "Original Value": float(row["original_value"]) if pd.notna(row["original_value"]) else None,
        "Original Unit": row["original_unit"],
        "Source Document": row["filename"],
        "Source Page": None if pd.isna(row["source_page"]) else int(row["source_page"]),
        "Extraction Method": row["extraction_method"],
        "Confidence": f"{float(row['confidence']) * 100:.1f}%",
        "Verification": row["verification_status"],
        "Source Snippet": row["source_snippet"],
    }


# ============================================================
# UNIQUE FEATURE 2: CONSISTENCY FIREWALL
# ============================================================


def consistency_firewall(record_audit: bool = False) -> dict[str, Any]:
    df = facts_df()
    docs = documents_df()
    findings = []
    errors = 0
    warnings = 0

    if df.empty:
        result = {"status": "NO_DATA", "errors": 0, "warnings": 0, "findings": []}
        if record_audit:
            audit("VALIDATION_RUN", "NO_DATA: no facts are indexed")
        return result

    # Facts in a group must agree after unit normalization.
    for key, grp in df.groupby(["entity", "metric", "period"], dropna=False):
        vals = grp["value"].round(6).unique().tolist()
        units = grp["unit"].fillna("").str.lower().unique().tolist()
        if len(vals) > 1 or len(units) > 1:
            errors += 1
            entity, metric, period = key
            detail = "; ".join(
                f"{r['filename']} = {r['value']:,.4f} {r['unit']}" for _, r in grp.iterrows()
            )
            findings.append(
                {
                    "severity": "ERROR",
                    "type": "SOURCE_CONFLICT",
                    "message": f"{entity} / {metric} / {period}",
                    "detail": detail,
                }
            )

    for _, row in df[df["value"] < 0].iterrows():
        errors += 1
        findings.append(
            {
                "severity": "ERROR",
                "type": "INVALID_NEGATIVE_VALUE",
                "message": f"{row['fact_id']} has a negative {row['metric']} value",
                "detail": f"{row['filename']} | {row['source_snippet']}",
            }
        )

    # Low confidence and missing-period facts require human review.
    low = df[df["confidence"] < 0.85]
    for _, r in low.iterrows():
        warnings += 1
        findings.append(
            {
                "severity": "WARNING",
                "type": "LOW_CONFIDENCE",
                "message": f"{r['fact_id']} confidence {r['confidence'] * 100:.1f}%",
                "detail": f"{r['filename']} | page {r['source_page']}",
            }
        )
    unknown_period = df[df["period"] == "Unknown Period"]
    for _, r in unknown_period.iterrows():
        warnings += 1
        findings.append(
            {
                "severity": "WARNING",
                "type": "UNKNOWN_PERIOD",
                "message": f"{r['fact_id']} has no confidently detected reporting period",
                "detail": f"{r['filename']} | {r['source_snippet']}",
            }
        )
    for _, r in df[df["period"] != "Unknown Period"].iterrows():
        match = re.fullmatch(r"FY(\d{4})-(\d{2})", str(r["period"]))
        if match and int(match.group(2)) != (int(match.group(1)) + 1) % 100:
            warnings += 1
            findings.append(
                {
                    "severity": "WARNING",
                    "type": "PERIOD_SANITY",
                    "message": f"{r['fact_id']} has an unusual financial-year range",
                    "detail": f"{r['period']} | {r['filename']}",
                }
            )

    # Recalculate achievement per source document so unrelated records cannot mix.
    document_content = docs.set_index("id")["content"].to_dict() if not docs.empty else {}
    for key, grp in df.groupby(["entity", "period", "document_id"], dropna=False):
        production = grp[grp["metric"] == "production"]
        target = grp[grp["metric"] == "target"]
        if not production.empty and not target.empty:
            p = float(production.iloc[0]["value"])
            t = float(target.iloc[0]["value"])
            if t != 0:
                achievement = p / t * 100
                if achievement < 0 or achievement > 150:
                    warnings += 1
                    findings.append(
                        {
                            "severity": "WARNING",
                            "type": "ARITHMETIC_SANITY",
                            "message": f"Unusual achievement = {achievement:.2f}%",
                            "detail": f"{key[0]} / {key[1]} | {grp.iloc[0]['filename']}",
                        }
                    )
                content = str(document_content.get(key[2], ""))
                stated = re.search(
                    r"\b(?:achievement(?:\s+rate)?|target\s+achieved)"
                    r"\D{0,20}([\d,]+(?:\.\d+)?)\s*%",
                    content,
                    re.I,
                )
                if stated:
                    stated_value = clean_number(stated.group(1))
                    if abs(stated_value - achievement) > max(1.0, abs(achievement) * 0.01):
                        errors += 1
                        findings.append(
                            {
                                "severity": "ERROR",
                                "type": "ARITHMETIC_INCONSISTENCY",
                                "message": (
                                    f"Reported achievement {stated_value:.2f}% does not match "
                                    f"recalculated {achievement:.2f}%"
                                ),
                                "detail": (f"{grp.iloc[0]['filename']} | {key[0]} / {key[1]}"),
                            }
                        )

    status = "PASS" if errors == 0 else "REVIEW_REQUIRED"
    result = {"status": status, "errors": errors, "warnings": warnings, "findings": findings}
    if record_audit:
        audit(
            "VALIDATION_RUN",
            f"{status}: {errors} error(s), {warnings} warning(s)",
        )
    return result


# ============================================================
# UNIQUE FEATURE 3: CHANGE INTELLIGENCE
# ============================================================


def compare_entity_periods(entity: str, old_period: str, new_period: str):
    df = facts_df()
    df = df[df["entity"].str.lower() == entity.lower()]

    old = df[df["period"] == old_period].copy()
    new = df[df["period"] == new_period].copy()

    metrics = sorted(set(old["metric"]).union(set(new["metric"])))
    rows = []

    for metric in metrics:
        o = old[old["metric"] == metric].sort_values("id", ascending=False).head(1)
        n = new[new["metric"] == metric].sort_values("id", ascending=False).head(1)

        ov = float(o.iloc[0]["value"]) if not o.empty else np.nan
        nv = float(n.iloc[0]["value"]) if not n.empty else np.nan
        unit = (
            str(n.iloc[0]["unit"]) if not n.empty else str(o.iloc[0]["unit"]) if not o.empty else ""
        )

        pct_change = np.nan
        absolute_change = np.nan
        if pd.notna(ov) and ov != 0 and pd.notna(nv):
            absolute_change = nv - ov
            pct_change = (nv - ov) / ov * 100
        elif pd.notna(ov) and pd.notna(nv):
            absolute_change = nv - ov

        if pd.isna(ov):
            status = "NEW"
        elif pd.isna(nv):
            status = "MISSING"
        elif math.isclose(ov, nv, rel_tol=1e-9, abs_tol=1e-9):
            status = "UNCHANGED"
        else:
            status = "CHANGED"

        risk = "NORMAL"
        if pd.notna(pct_change) and abs(pct_change) >= 20:
            risk = "ATTENTION"

        rows.append(
            {
                "Metric": metric,
                old_period: ov,
                new_period: nv,
                "Absolute Change": absolute_change,
                "% Change": pct_change,
                "Unit": unit,
                "Status": status,
                "Attention": risk,
            }
        )

    return pd.DataFrame(rows)


# ============================================================
# SEARCH / EVIDENCE RETRIEVAL
# ============================================================


def retrieve_sources(query: str, k: int = 5):
    docs = documents_df()
    if docs.empty:
        return []

    corpus = docs["content"].fillna("").tolist()

    if TfidfVectorizer is None:
        terms = [x.lower() for x in re.findall(r"\w+", query) if len(x) > 2]
        scored = []
        for _, row in docs.iterrows():
            text = str(row["content"]).lower()
            score = sum(text.count(t) for t in terms)
            scored.append((score, row.to_dict()))
        scored.sort(key=lambda x: x[0], reverse=True)
        return [r for score, r in scored[:k] if score > 0]

    vec = TfidfVectorizer(stop_words="english", ngram_range=(1, 2))
    mat = vec.fit_transform(corpus)
    q = vec.transform([query])
    sims = cosine_similarity(q, mat)[0]
    order = np.argsort(sims)[::-1][:k]

    out = []
    for idx in order:
        if sims[idx] <= 0:
            continue
        row = docs.iloc[idx].to_dict()
        row["relevance"] = float(sims[idx])
        out.append(row)
    return out


def detect_query_metric(query: str) -> str | None:
    q = query.lower()
    aliases = sorted(
        ((key, alias) for key, meta in ONTOLOGY.items() for alias in meta["aliases"]),
        key=lambda item: len(item[1]),
        reverse=True,
    )
    for key, alias in aliases:
        found = (
            re.search(rf"\b{re.escape(alias)}\b", q)
            if len(alias) <= 3
            else re.search(re.escape(alias), q)
        )
        if found:
            return key
    return None


def find_best_fact(query: str):
    df = facts_df()
    if df.empty:
        return None

    q = query.lower()
    metric = detect_query_metric(query)

    period = detect_period(query)

    subset = df.copy()
    if metric:
        subset = subset[subset["metric"] == metric]
    if period != "Unknown Period":
        subset = subset[subset["period"] == period]

    names = set()
    for column in ("entity", "mine", "project", "subsidiary"):
        if column in df.columns:
            names.update(
                str(value)
                for value in df[column].dropna().unique()
                if str(value).strip().lower() not in {"unknown", "unknown entity"}
            )
    matched_names = [
        name for name in names if re.search(rf"(?<!\w){re.escape(name.lower())}(?!\w)", q)
    ]
    if matched_names:
        subset = subset[
            subset.apply(
                lambda row: any(
                    str(row.get(column, "")).lower() == name.lower()
                    for column in ("entity", "mine", "project", "subsidiary")
                    for name in matched_names
                ),
                axis=1,
            )
        ]
    elif (
        re.search(r"\b(?:mine|project|subsidiary|block)\s+[A-Za-z0-9]", query, re.I)
        or len(subset) != 1
    ):
        return None

    if subset.empty:
        return None
    return subset.sort_values(["confidence", "id"], ascending=[False, False]).iloc[0]


def evidence_answer(query: str):
    df = facts_df()
    best = None
    matched_facts = []
    comparisons = []
    requested_periods = [
        normalize_period(match.group(0))
        for match in re.finditer(r"(?:FY\s*)?20\d{2}\s*[-–/]\s*(?:\d{2}|20\d{2})", query, re.I)
    ]
    if any(term in query.lower() for term in ("highest", "maximum", "most")):
        metric = detect_query_metric(query) or "production"
        candidates = df[df["metric"] == metric] if not df.empty else df
        if not candidates.empty:
            best = candidates.loc[candidates["value"].idxmax()]
            matched_facts = [best]
    elif len(requested_periods) >= 2 and not df.empty:
        best = find_best_fact(query)
        if best is not None:
            matched_facts = [
                row
                for _, row in df[
                    (df["entity"] == best["entity"])
                    & (df["metric"] == best["metric"])
                    & (df["period"].isin(requested_periods))
                ]
                .sort_values("period")
                .iterrows()
            ]
            if len(matched_facts) >= 2:
                old_fact, new_fact = matched_facts[0], matched_facts[-1]
                comparisons.append(build_comparison(old_fact, new_fact))
            else:
                best = None
                matched_facts = []
        else:
            metric = detect_query_metric(query)
            entity_was_requested = bool(
                re.search(r"\b(?:mine|project|subsidiary|block)\s+[A-Za-z0-9]", query, re.I)
            )
            if metric and not entity_was_requested:
                candidates = df[(df["metric"] == metric) & (df["period"].isin(requested_periods))]
                for _, group in candidates.groupby("entity"):
                    by_period = {
                        period: group[group["period"] == period]
                        .sort_values("id", ascending=False)
                        .head(1)
                        for period in requested_periods
                    }
                    if all(not period_facts.empty for period_facts in by_period.values()):
                        old_fact = by_period[requested_periods[0]].iloc[0]
                        new_fact = by_period[requested_periods[-1]].iloc[0]
                        comparisons.append(build_comparison(old_fact, new_fact))
                        matched_facts.extend([old_fact, new_fact])
                if comparisons:
                    best = matched_facts[-1]
    else:
        best = find_best_fact(query)
        if best is not None:
            matched_facts = [best]

    sources = retrieve_sources(query, 5)

    if best is None and not sources:
        return {
            "answer": "No evidence is indexed yet. Upload approved demo/source documents.",
            "fact": None,
            "facts": [],
            "comparisons": [],
            "sources": [],
        }

    if best is None:
        answer = (
            (
                "Relevant source documents were found, but no unambiguous structured fact "
                "matched the requested metric, entity, and period. No numerical value is "
                "provided without a matching Fact Passport."
            )
            if sources
            else (
                "No matching fact or source document is indexed. No numerical value was generated."
            )
        )
    else:
        if comparisons:
            answer_rows = []
            for pair in comparisons:
                pct_text = (
                    f"{pair['percentage_change']:+.2f}%"
                    if pair["percentage_change"] is not None
                    else "percentage change unavailable"
                )
                citations = ", ".join(
                    fact_id for fact_id in (pair["old_fact_id"], pair["new_fact_id"])
                )
                answer_rows.append(
                    f"**{pair['entity']}**: {pair['old_value']:,.4f} "
                    f"{pair['unit']} ({pair['old_period']}) to "
                    f"{pair['new_value']:,.4f} {pair['unit']} ({pair['new_period']}); "
                    f"absolute change {pair['absolute_change']:+,.4f} {pair['unit']} "
                    f"({pct_text}). Fact IDs: {citations}."
                )
            answer = f"**{best['metric'].replace('_', ' ').title()} changes:**\n\n"
            answer += "\n\n".join(answer_rows)
        else:
            value = f"{float(best['value']):,.4f}".rstrip("0").rstrip(".")
            if any(term in query.lower() for term in ("highest", "maximum", "most")):
                answer = (
                    f"The highest indexed **{best['metric'].replace('_', ' ').title()}** "
                    f"is **{value} {best['unit']}** for **{best['entity']}** "
                    f"in **{best['period']}**."
                )
            else:
                answer = (
                    f"**{best['metric'].replace('_', ' ').title()}** for "
                    f"**{best['entity']}** in **{best['period']}** is "
                    f"**{value} {best['unit']}**."
                )
            page_text = (
                f", page **{int(best['source_page'])}**"
                if pd.notna(best["source_page"])
                else " (page not recorded)"
            )
            answer += (
                f" Fact ID: **{best['fact_id']}**. Source document: "
                f"**{best['filename']}**{page_text}."
            )
        answer += f" Confidence: **{float(best['confidence']) * 100:.1f}%**."

    return {
        "answer": answer,
        "fact": best,
        "facts": matched_facts,
        "comparisons": comparisons,
        "sources": sources,
    }


def build_comparison(old_fact: pd.Series, new_fact: pd.Series) -> dict[str, Any]:
    old_value = float(old_fact["value"])
    new_value = float(new_fact["value"])
    absolute_change = new_value - old_value
    return {
        "entity": str(old_fact["entity"]),
        "metric": str(old_fact["metric"]),
        "old_period": str(old_fact["period"]),
        "new_period": str(new_fact["period"]),
        "old_value": old_value,
        "new_value": new_value,
        "unit": str(new_fact["unit"]),
        "absolute_change": absolute_change,
        "percentage_change": (absolute_change / old_value * 100 if old_value != 0 else None),
        "old_fact_id": str(old_fact["fact_id"]),
        "new_fact_id": str(new_fact["fact_id"]),
        "old_source": str(old_fact["filename"]),
        "new_source": str(new_fact["filename"]),
        "old_page": (int(old_fact["source_page"]) if pd.notna(old_fact["source_page"]) else None),
        "new_page": (int(new_fact["source_page"]) if pd.notna(new_fact["source_page"]) else None),
    }


# ============================================================
# UNIQUE FEATURE 4: PARLIAMENTARY QUESTION COPILOT
# ============================================================


def analyze_question(question: str):
    q = question.lower()
    metrics = [
        key
        for key, meta in ONTOLOGY.items()
        if any(
            bool(
                re.search(rf"\b{re.escape(alias)}\b", q)
                if len(alias) <= 3
                else re.search(re.escape(alias), q)
            )
            for alias in meta["aliases"]
        )
    ]

    years = re.findall(r"(?:FY\s*)?20\d{2}\s*[-–/]\s*(?:\d{2}|20\d{2})", question, re.I)
    requires_growth = any(x in q for x in ["growth", "increase", "decrease", "change", "trend"])
    requires_comparison = any(
        x in q for x in ["compare", "comparison", "last", "previous", "historical"]
    )
    locations = []
    for word in [
        "state-wise",
        "statewise",
        "district-wise",
        "subsidiary-wise",
        "mine-wise",
        "project-wise",
    ]:
        if word in q:
            locations.append(word)
    indexed_facts = facts_df()
    entities = sorted(
        {
            str(value)
            for column in ("entity", "mine", "subsidiary", "project")
            for value in indexed_facts.get(column, pd.Series(dtype=str)).dropna().unique()
            if str(value).lower() not in {"unknown", "unknown entity"}
            and re.search(rf"(?<!\w){re.escape(str(value).lower())}(?!\w)", q)
        }
    )

    return {
        "metrics": metrics,
        "periods": [normalize_period(x) for x in years],
        "requires_growth": requires_growth,
        "requires_comparison": requires_comparison,
        "dimensions_detected": locations,
        "entities": entities,
    }


def draft_parliamentary_answer(question: str):
    analysis = analyze_question(question)
    result = evidence_answer(question)

    draft = (
        "### Draft Response\n\n"
        "This is an AI-assisted draft prepared from indexed evidence. "
        "It requires authorized human review before official use.\n\n"
    )

    if result["fact"] is not None and result.get("facts"):
        f = result["fact"]
        draft += f"Based on indexed records, {f['metric'].replace('_', ' ')}:\n"
        for pair in result.get("comparisons", []):
            pct = (
                f"{pair['percentage_change']:.2f}%"
                if pair["percentage_change"] is not None
                else "unavailable"
            )
            draft += (
                f"- {pair['entity']}: {pair['old_value']:,.4f} {pair['unit']} "
                f"({pair['old_period']}) to {pair['new_value']:,.4f} {pair['unit']} "
                f"({pair['new_period']}); absolute change "
                f"{pair['absolute_change']:+,.4f} {pair['unit']}; "
                f"percentage change {pct}. Fact IDs: {pair['old_fact_id']}, "
                f"{pair['new_fact_id']}; sources: {pair['old_source']} "
                f"(page {pair['old_page']}), {pair['new_source']} "
                f"(page {pair['new_page']}).\n"
            )
        if not result.get("comparisons"):
            for evidence in result["facts"]:
                draft += (
                    f"- {evidence['period']}: {float(evidence['value']):,.4f} "
                    f"{evidence['unit']} [Fact ID {evidence['fact_id']}; "
                    f"source {evidence['filename']}, page {evidence['source_page']}].\n"
                )
        if not result.get("comparisons") and len(result["facts"]) > 1:
            ordered = sorted(result["facts"], key=lambda fact: fact["period"])
            old, new = ordered[0], ordered[-1]
            difference = float(new["value"]) - float(old["value"])
            draft += f"- Deterministic absolute change: {difference:,.4f} {new['unit']}."
            if float(old["value"]) != 0:
                draft += f" Percentage change: {difference / float(old['value']) * 100:.2f}%."
            draft += "\n\n"
        else:
            draft += "\n"
    else:
        draft += "No structured figure could be confidently extracted for the exact question.\n\n"

    draft += "**Detected question requirements:**\n"
    draft += f"- Metrics: {', '.join(analysis['metrics']) or 'Not confidently detected'}\n"
    draft += f"- Periods: {', '.join(analysis['periods']) or 'Not explicitly stated'}\n"
    draft += f"- Comparison needed: {'Yes' if analysis['requires_comparison'] else 'No'}\n"
    draft += f"- Growth/change calculation: {'Yes' if analysis['requires_growth'] else 'No'}\n"
    draft += f"- Dimensions: {', '.join(analysis['dimensions_detected']) or 'None detected'}\n"
    draft += f"- Entities: {', '.join(analysis['entities']) or 'None confidently matched'}\n"
    draft += "\n**DRAFT — requires authorized human approval before official use.**\n"

    return analysis, draft, result


# ============================================================
# UNIQUE FEATURE 5: TOPIC EVOLUTION
# ============================================================


def top_keywords(text: str, n=30):
    tokens = re.findall(r"[A-Za-z][A-Za-z0-9-]{2,}", text.lower())
    freq = {}
    for t in tokens:
        if t in STOPWORDS or t.isdigit():
            continue
        freq[t] = freq.get(t, 0) + 1
    return sorted(freq.items(), key=lambda x: x[1], reverse=True)[:n]


def domain_topic_scores(text: str):
    t = text.lower()
    groups = {
        "Production": ["production", "output", "dispatch", "target"],
        "Geology & Exploration": [
            "geology",
            "geological",
            "borehole",
            "drilling",
            "seam",
            "reserve",
        ],
        "Mining Operations": ["mining", "overburden", "excavator", "equipment", "stripping"],
        "Coal Quality": ["gcv", "grade", "ash", "moisture", "quality"],
        "Environment": ["environment", "water", "dust", "emission", "reclamation"],
        "Safety": ["safety", "accident", "incident", "risk"],
        "Projects & Land": ["land", "project", "clearance", "acquisition"],
    }
    scores = {topic: sum(t.count(term) for term in words) for topic, words in groups.items()}
    return dict(sorted(scores.items(), key=lambda x: x[1], reverse=True))


# ============================================================
# REPORT GENERATION
# ============================================================


def generate_docx(
    title: str,
    selected: pd.DataFrame,
    firewall: dict,
    comparison: pd.DataFrame | None = None,
    approval_note: str = "DRAFT - requires authorized human approval before official use.",
):
    if Document is None:
        raise RuntimeError("Install python-docx first.")

    doc = Document()
    doc.add_heading(title, 0)
    doc.add_paragraph(f"{approval_note} Generated on {datetime.now().strftime('%d-%m-%Y %H:%M')}.")

    doc.add_heading("1. Executive Summary", level=1)
    doc.add_paragraph(
        "This report was generated from indexed CMPDI/CIL-style records using "
        "structured extraction, evidence provenance and automated validation. "
        "All included figures are linked to their Fact Passport evidence."
    )

    doc.add_heading("2. Extracted Facts", level=1)
    if selected.empty:
        doc.add_paragraph("No facts selected.")
    else:
        table = doc.add_table(rows=1, cols=8)
        headers = ["Fact ID", "Entity", "Metric", "Period", "Value", "Unit", "Confidence", "Source"]
        for i, h in enumerate(headers):
            table.rows[0].cells[i].text = h

        for _, r in selected.head(50).iterrows():
            cells = table.add_row().cells
            values = [
                r["fact_id"],
                r["entity"],
                r["metric"],
                r["period"],
                f"{float(r['value']):,.4f}".rstrip("0").rstrip("."),
                r["unit"],
                f"{float(r['confidence']) * 100:.1f}%",
                f"{r['filename']} / p.{r['source_page']}",
            ]
            for i, v in enumerate(values):
                cells[i].text = str(v)
            doc.add_paragraph(
                f"Evidence {r['fact_id']}: {r['source_snippet']} "
                f"(source: {r['filename']}, page {r['source_page']})."
            )

    doc.add_heading("3. Period Comparisons", level=1)
    if comparison is None or comparison.empty:
        doc.add_paragraph("No period comparison was selected.")
    else:
        comparison_table = doc.add_table(rows=1, cols=len(comparison.columns))
        for index, column in enumerate(comparison.columns):
            comparison_table.rows[0].cells[index].text = str(column)
        for _, comparison_row in comparison.iterrows():
            cells = comparison_table.add_row().cells
            for index, value in enumerate(comparison_row):
                cells[index].text = str(value)

    doc.add_heading("4. Consistency Firewall", level=1)
    doc.add_paragraph(
        f"Status: {firewall['status']} | "
        f"Errors: {firewall['errors']} | Warnings: {firewall['warnings']}"
    )
    for item in firewall["findings"][:30]:
        doc.add_paragraph(
            f"{item['severity']} | {item['type']} | {item['message']} | {item['detail']}"
        )

    doc.add_heading("5. Approval Note", level=1)
    doc.add_paragraph(approval_note)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def generate_pdf(
    title: str,
    selected: pd.DataFrame,
    firewall: dict,
    comparison: pd.DataFrame | None = None,
    approval_note: str = "DRAFT - requires authorized human approval before official use.",
):
    if SimpleDocTemplate is None:
        raise RuntimeError("Install reportlab first.")

    buf = io.BytesIO()
    pdf = SimpleDocTemplate(
        buf, pagesize=A4, rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30
    )
    styles = getSampleStyleSheet()

    def pdf_text(value: Any) -> str:
        return str(value).encode("latin-1", errors="replace").decode("latin-1")

    story = [
        Paragraph(html.escape(pdf_text(title)), styles["Title"]),
        Spacer(1, 10),
        Paragraph(
            html.escape(
                pdf_text(f"{approval_note} Generated {datetime.now().strftime('%d-%m-%Y %H:%M')}.")
            ),
            styles["BodyText"],
        ),
        Spacer(1, 10),
        Paragraph("Consistency Firewall", styles["Heading2"]),
        Paragraph(
            f"Status: {firewall['status']} | Errors: {firewall['errors']} | Warnings: {firewall['warnings']}",
            styles["BodyText"],
        ),
        Spacer(1, 10),
        Paragraph("Executive Summary", styles["Heading2"]),
        Paragraph(
            f"{len(selected)} sourced fact(s) included. {html.escape(pdf_text(approval_note))}",
            styles["BodyText"],
        ),
        Spacer(1, 10),
        Paragraph("Extracted Figures", styles["Heading2"]),
    ]

    data = [["Fact", "Entity", "Metric", "Period", "Value", "Unit", "Conf."]]
    for _, r in selected.head(30).iterrows():
        data.append(
            [
                pdf_text(r["fact_id"]),
                pdf_text(r["entity"])[:25],
                pdf_text(r["metric"]),
                pdf_text(r["period"]),
                f"{float(r['value']):.4f}".rstrip("0").rstrip("."),
                pdf_text(r["unit"]),
                f"{float(r['confidence']) * 100:.0f}%",
            ]
        )

    if len(data) == 1:
        data.append(["-", "-", "-", "-", "-", "-", "-"])

    table = Table(data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#274c77")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                ("FONTSIZE", (0, 0), (-1, -1), 7),
            ]
        )
    )
    story.extend([table, Spacer(1, 10)])
    story.append(Paragraph("Period Comparisons", styles["Heading2"]))
    if comparison is None or comparison.empty:
        story.append(Paragraph("No period comparison was selected.", styles["BodyText"]))
    else:
        comparison_data = [[pdf_text(column) for column in comparison.columns]] + [
            [pdf_text(value) for value in row]
            for row in comparison.itertuples(index=False, name=None)
        ]
        comparison_table = Table(comparison_data, repeatRows=1)
        comparison_table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#274c77")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                    ("FONTSIZE", (0, 0), (-1, -1), 7),
                ]
            )
        )
        story.extend([comparison_table, Spacer(1, 10)])
    story.append(Paragraph("Source Evidence", styles["Heading2"]))
    for _, row in selected.head(30).iterrows():
        evidence = (
            f"{row['fact_id']} | {row['filename']} | page {row['source_page']} | "
            f"{row['source_snippet']}"
        )
        story.append(Paragraph(html.escape(pdf_text(evidence)), styles["BodyText"]))
        story.append(Spacer(1, 4))
    story.append(Paragraph("Firewall Findings", styles["Heading2"]))
    for item in firewall["findings"][:30]:
        finding = f"{item['severity']} | {item['type']} | {item['message']} | {item['detail']}"
        story.append(Paragraph(html.escape(pdf_text(finding)), styles["BodyText"]))
    story.append(Paragraph(html.escape(pdf_text(approval_note)), styles["BodyText"]))

    pdf.build(story)
    return buf.getvalue()


# ============================================================
# OPTIONAL OLLAMA
# ============================================================


def ollama_rewrite(prompt: str, model: str, host: str = "http://localhost:11434"):
    if requests is None:
        raise RuntimeError("The requests package is required for the optional Ollama feature.")
    if not model.strip():
        raise ValueError("Enter the name of a locally installed Ollama model.")
    try:
        r = requests.post(
            f"{host.rstrip('/')}/api/generate",
            json={"model": model, "prompt": prompt, "stream": False},
            timeout=90,
        )
        r.raise_for_status()
        response = r.json().get("response")
        if not isinstance(response, str) or not response.strip():
            raise RuntimeError("Ollama returned an empty wording response.")
        return response
    except requests.RequestException as exc:
        raise RuntimeError(
            f"Could not reach Ollama at {host}. Verify Ollama is running and the model is available."
        ) from exc
    except ValueError as exc:
        raise RuntimeError("Ollama returned an invalid JSON response.") from exc


# ============================================================
# STREAMLIT UI
# ============================================================


def render_visual_system(active_page: str):
    module_content = {
        "Dashboard": (
            "⛏️",
            "Your operations, at a glance",
            "A live view of indexed evidence, verification, and reporting activity.",
            0,
        ),
        "Document Center": (
            "📄",
            "Bring records into focus",
            "Upload and extract source material while keeping every page and table in view.",
            0,
        ),
        "Fact Passport": (
            "🪪",
            "Every figure has a lineage",
            "Inspect the source, context, confidence, and verification status behind a fact.",
            1,
        ),
        "Consistency Firewall": (
            "🛡️",
            "Confidence before circulation",
            "Surface conflicts and arithmetic anomalies before they reach a final report.",
            2,
        ),
        "Change Intelligence": (
            "📈",
            "Read the change, not just the number",
            "Compare reporting periods and identify the changes that deserve attention.",
            2,
        ),
        "AI Query Assistant": (
            "💬",
            "Ask the indexed evidence",
            "Get grounded answers with calculations performed from structured facts.",
            3,
        ),
        "Parliamentary Copilot": (
            "🏛️",
            "Draft with evidence in hand",
            "Turn priority questions into sourced drafts for authorized human review.",
            3,
        ),
        "Topic Intelligence": (
            "☁️",
            "See the themes in your records",
            "Explore mining topics, keywords, and how attention shifts across periods.",
            3,
        ),
        "Report Generator": (
            "📝",
            "From verified facts to a polished report",
            "Build traceable outputs, then route them through the human approval workflow.",
            4,
        ),
        "Mining Ontology": (
            "🧭",
            "Speak the language of mining",
            "Tune domain terms, aliases, and units to fit your reporting vocabulary.",
            1,
        ),
        "Audit Log": (
            "🔎",
            "A clear record of every decision",
            "Review document, validation, verification, and report activity.",
            4,
        ),
    }
    icon, heading, description, active_step = module_content[active_page]
    st.markdown(
        """
        <style>

        :root {
            color-scheme: dark;
            --geo-canvas: #0b1218;
            --geo-panel: #111c26;
            --geo-card: rgba(20, 32, 43, 0.88);
            --geo-card-hover: rgba(27, 43, 56, 0.96);
            --geo-border: rgba(189, 207, 220, 0.1);
            --geo-border-glow: rgba(0, 240, 181, 0.32);
            --geo-text-head: #f3f7fa;
            --geo-text-body: #c2d1dc;
            --geo-text-muted: #7e94a4;
            --geo-emerald: #00f0b5;
            --geo-cyan: #00d8f6;
            --geo-copper: #f5ba6b;
            --geo-crimson: #ff4d5e;
            --geo-amber: #ffb238;
            --geo-shadow-card:
                0 12px 32px -4px rgba(2, 9, 14, 0.32),
                0 4px 12px -2px rgba(2, 9, 14, 0.2);
            --geo-shadow-hover:
                0 20px 40px -6px rgba(0, 240, 181, 0.12),
                0 8px 24px -4px rgba(0, 0, 0, 0.45);
            --geo-radius-sm: 10px;
            --geo-radius-md: 14px;
            --geo-radius-lg: 20px;
            --geo-ease: cubic-bezier(0.16, 1, 0.3, 1);
        }

        html, body, [class*="css"], .stApp {
            font-family: Inter, "Segoe UI", -apple-system, BlinkMacSystemFont, sans-serif !important;
            background-color: var(--geo-canvas) !important;
            color: var(--geo-text-body) !important;
            -webkit-font-smoothing: antialiased;
            text-rendering: optimizeLegibility;
        }
        ::selection { background: rgba(0,240,181,.26); color: #f3f7fa; }
        * { scrollbar-color: #354b59 #111c26; scrollbar-width: thin; }
        *::-webkit-scrollbar { width: 9px; height: 9px; }
        *::-webkit-scrollbar-track { background: #111c26; }
        *::-webkit-scrollbar-thumb {
            border: 2px solid #111c26;
            border-radius: 999px;
            background: #354b59;
        }
        *::-webkit-scrollbar-thumb:hover { background: #537080; }
        [data-testid="stAppViewContainer"] { background: transparent !important; }
        [data-testid="stAppViewContainer"] section [data-testid="stVerticalBlock"] {
            gap: 1rem;
        }
        [data-testid="stMarkdownContainer"] { color: var(--geo-text-body); }
        [data-testid="stMarkdownContainer"] h1,
        [data-testid="stMarkdownContainer"] h2,
        [data-testid="stMarkdownContainer"] h3,
        [data-testid="stMarkdownContainer"] h4 {
            color: var(--geo-text-head);
            font-weight: 700;
            letter-spacing: -0.035em;
        }
        [data-testid="stMarkdownContainer"] h2 {
            margin-top: 1.45rem;
            margin-bottom: 0.7rem;
            font-size: clamp(1.25rem, 2vw, 1.55rem);
        }
        [data-testid="stMarkdownContainer"] h3 { font-size: 1.12rem; }
        [data-testid="stMarkdownContainer"] h4 { font-size: 0.98rem; }
        [data-testid="stCaptionContainer"] p { color: var(--geo-text-muted); }

        .stApp {
            background:
                radial-gradient(ellipse 90% 50% at 85% -5%, rgba(0, 240, 181, 0.075), transparent 60%),
                radial-gradient(ellipse 70% 40% at 5% 25%, rgba(0, 216, 246, 0.055), transparent 55%),
                linear-gradient(175deg, #0b1218 0%, #0e1720 48%, #121d27 100%);
            background-size: 180% 180%, 170% 170%, 100% 100%;
            background-position: 0% 0%, 100% 100%, 0% 0%;
            animation: ambient-drift 48s ease-in-out infinite alternate;
        }

        /* Top Header */
        [data-testid="stHeader"] {
            background: rgba(11, 18, 24, 0.84) !important;
            backdrop-filter: blur(20px) !important;
            border-bottom: 1px solid var(--geo-border) !important;
        }
        [data-testid="stToolbar"] { right: 1rem; }
        [data-testid="stDecoration"] { display: none; }

        [data-testid="stMainBlockContainer"] {
            max-width: 1540px !important;
            padding: 1.5rem clamp(1rem, 2.5vw, 2.8rem) 3.5rem !important;
        }

        /* Sidebar Styling */
        [data-testid="stSidebar"] {
            background: linear-gradient(180deg, #111c26 0%, #0b1218 100%) !important;
            border-right: 1px solid var(--geo-border) !important;
            box-shadow: 12px 0 35px rgba(0, 0, 0, 0.4) !important;
        }
        [data-testid="stSidebar"] > div:first-child {
            border-right: 1px solid rgba(255,255,255,.035);
        }
        [data-testid="stSidebar"] [data-testid="stMarkdownContainer"] p {
            color: var(--geo-text-muted);
        }
        [data-testid="stSidebar"] button[kind="header"] {
            color: var(--geo-text-muted);
            transition: color .2s ease, background .2s ease;
        }
        [data-testid="stSidebar"] button[kind="header"]:hover {
            color: var(--geo-emerald);
            background: rgba(0,240,181,.08);
        }

        .sidebar-brand {
            padding: 1.1rem 1rem;
            margin-bottom: 1.25rem;
            border: 1px solid rgba(0, 240, 181, 0.18);
            border-radius: 18px;
            background: linear-gradient(145deg, rgba(20, 32, 44, 0.95), rgba(12, 20, 28, 0.9));
            box-shadow: 0 10px 24px rgba(0, 0, 0, 0.3), inset 0 1px 0 rgba(255, 255, 255, 0.06);
        }
        .sidebar-brand-mark {
            display: inline-grid; place-items: center; width: 40px; height: 40px;
            margin-bottom: 0.75rem; border-radius: 12px;
            background: linear-gradient(135deg, #00F0B5, #0099B8);
            color: #041014; font-size: 1.25rem; font-weight: 800;
            box-shadow: 0 4px 14px rgba(0, 240, 181, 0.35);
        }
        .sidebar-brand-title {
            color: var(--geo-text-head); font-size: 1rem; font-weight: 800;
            letter-spacing: -0.025em;
        }
        .sidebar-brand-caption {
            margin-top: 0.3rem; color: var(--geo-text-muted); font-size: 0.72rem;
            line-height: 1.4;
        }

        [data-testid="stSidebar"] [data-testid="stRadio"] label {
            border-radius: 12px;
            padding: 0.42rem 0.65rem;
            margin: 0.1rem 0;
            font-size: 0.85rem;
            font-weight: 500;
            transition: all 0.25s cubic-bezier(0.16, 1, 0.3, 1);
        }
        [data-testid="stSidebar"] [data-testid="stRadio"] label:hover {
            background: rgba(0, 240, 181, 0.08);
            color: #FFFFFF !important;
            transform: translateX(4px);
        }
        [data-testid="stSidebar"] [data-testid="stRadio"] label:has(input:checked) {
            border: 1px solid rgba(0,240,181,.2);
            background: linear-gradient(100deg, rgba(0,240,181,.13), rgba(0,216,246,.045));
            box-shadow: inset 2px 0 0 var(--geo-emerald);
            color: #fff !important;
        }

        /* Hero Banner */
        .mine-hero {
            isolation: isolate;
            position: relative;
            display: flex;
            align-items: center;
            justify-content: space-between;
            gap: 2rem;
            min-height: 230px;
            margin: 0.25rem 0 1.5rem;
            padding: clamp(1.8rem, 3.5vw, 2.8rem);
            overflow: hidden;
            border: 1px solid rgba(0, 240, 181, 0.2);
            border-radius: 24px;
            background:
                radial-gradient(ellipse at 80% 40%, rgba(0, 240, 181, 0.12), transparent 45%),
                radial-gradient(ellipse at 15% 90%, rgba(0, 216, 246, 0.08), transparent 40%),
                linear-gradient(135deg, rgba(16, 26, 36, 0.96) 0%, rgba(10, 17, 24, 0.94) 100%);
            background-size: 180% 180%, 170% 170%, 100% 100%;
            background-position: 0% 50%, 100% 50%, 0% 0%;
            box-shadow: var(--geo-shadow-card), inset 0 1px 0 rgba(255, 255, 255, 0.08);
            animation: hero-aura-drift 34s ease-in-out infinite alternate;
            transform-style: preserve-3d;
            perspective: 1200px;
        }
        .mine-hero-copy {
            position: relative;
            z-index: 2;
            transform: translateZ(24px);
        }
        .mine-eyebrow {
            display: inline-flex; align-items: center; gap: 0.55rem;
            margin-bottom: 0.85rem; color: var(--geo-emerald); font-size: 0.75rem;
            font-weight: 700; letter-spacing: 0.14em; text-transform: uppercase;
        }
        .mine-live-dot {
            width: 8px; height: 8px; border-radius: 50%;
            background: var(--geo-emerald); box-shadow: 0 0 0 4px rgba(0, 240, 181, 0.22);
            animation: pulse-soft 3.2s ease-in-out infinite;
        }
        .mine-hero h1 {
            margin: 0; color: var(--geo-text-head); font-size: clamp(2rem, 3.8vw, 3rem);
            font-weight: 800; line-height: 1.08; letter-spacing: -0.04em;
        }
        .mine-hero p {
            max-width: 650px; margin: 0.85rem 0 0; color: var(--geo-text-body);
            font-size: clamp(0.92rem, 1.2vw, 1.02rem); line-height: 1.6;
        }
        .mine-chip {
            padding: 0.35rem 0.75rem; border: 1px solid rgba(255, 255, 255, 0.09);
            border-radius: 999px; background: rgba(255, 255, 255, 0.04);
            color: var(--geo-text-head); font-size: 0.74rem; font-weight: 600;
            backdrop-filter: blur(8px);
        }

        /* Workflow Step Bar */
        .mine-workflow {
            display: grid; grid-template-columns: repeat(5, minmax(0, 1fr));
            gap: 0.85rem; margin: 0 0 2.2rem;
        }
        .mine-step {
            position: relative; padding: 0.9rem 1rem;
            border: 1px solid var(--geo-border); border-radius: 16px;
            background: var(--geo-card);
            box-shadow: 0 8px 20px rgba(0, 0, 0, 0.2);
            color: var(--geo-text-body); font-size: 0.82rem; font-weight: 600;
            backdrop-filter: blur(12px);
            transition: all 0.3s cubic-bezier(0.16, 1, 0.3, 1);
        }
        .mine-step.is-active {
            border-color: var(--geo-emerald);
            background: linear-gradient(145deg, rgba(0, 240, 181, 0.12), rgba(16, 26, 36, 0.9));
            box-shadow: 0 8px 24px rgba(0, 240, 181, 0.18), inset 0 0 0 1px rgba(0, 240, 181, 0.2);
        }
        .mine-step.is-active span { color: var(--geo-copper); }
        .mine-step:hover {
            transform: translateY(-3px);
            border-color: rgba(255, 255, 255, 0.16);
        }
        .mine-step span {
            display: block; margin-bottom: 0.25rem; color: var(--geo-emerald);
            font-size: 0.65rem; font-weight: 800; letter-spacing: 0.12em;
        }

        /* SaaS KPI Cards */
        .saas-kpi {
            position: relative; min-height: 146px; padding: 1.2rem 1.3rem 1.1rem;
            margin: 0 0 1.25rem 0;
            overflow: hidden; border: 1px solid var(--geo-border);
            border-radius: 20px;
            background: var(--geo-card);
            box-shadow: var(--geo-shadow-card);
            backdrop-filter: blur(14px);
            transition: all 0.35s cubic-bezier(0.16, 1, 0.3, 1);
        }
        .saas-kpi:hover {
            transform: translateY(-4px);
            border-color: var(--geo-border-glow);
            box-shadow: var(--geo-shadow-hover);
        }
        .saas-kpi-top {
            display: flex; align-items: center; justify-content: space-between;
        }
        .saas-kpi-label {
            color: var(--geo-text-muted); font-size: 0.78rem; font-weight: 700;
            letter-spacing: 0.04em; text-transform: uppercase;
        }
        .saas-kpi-icon {
            display: grid; place-items: center; width: 36px; height: 36px;
            border: 1px solid rgba(0, 240, 181, 0.24); border-radius: 12px;
            background: rgba(0, 240, 181, 0.08);
            color: var(--geo-emerald); font-size: 1rem;
        }
        .saas-kpi-value {
            font-family: ui-monospace, "Cascadia Code", Consolas, monospace;
            margin-top: 0.5rem; color: var(--geo-text-head);
            font-size: clamp(1.85rem, 2.6vw, 2.3rem);
            font-weight: 800; letter-spacing: -0.04em; line-height: 1.05;
        }
        .saas-kpi-detail {
            margin-top: 0.4rem; color: var(--geo-text-muted); font-size: 0.73rem;
            font-weight: 500;
        }
        .saas-kpi-accent {
            position: absolute; right: -20px; bottom: -30px; width: 100px; height: 75px;
            border: 1px solid rgba(0, 240, 181, 0.12); border-radius: 50%;
            transform: rotate(-20deg);
        }
        .saas-kpi-warning { border-color: rgba(255, 178, 56, 0.28); }
        .saas-kpi-warning .saas-kpi-icon {
            border-color: rgba(255, 178, 56, 0.35); background: rgba(255, 178, 56, 0.1);
            color: var(--geo-amber);
        }
        .saas-kpi-danger { border-color: rgba(255, 77, 94, 0.32); }
        .saas-kpi-danger .saas-kpi-icon {
            border-color: rgba(255, 77, 94, 0.38); background: rgba(255, 77, 94, 0.1);
            color: var(--geo-crimson);
        }

        /* Buttons & CTA */
        .stButton > button, .stDownloadButton > button {
            border: 1px solid rgba(0, 240, 181, 0.35) !important;
            border-radius: 12px !important;
            background: linear-gradient(135deg, #00F0B5 0%, #00C49F 50%, #0099B8 100%) !important;
            color: #041418 !important;
            font-weight: 750 !important;
            font-size: 0.9rem !important;
            padding: 0.6rem 1.4rem !important;
            box-shadow: 0 4px 16px rgba(0, 240, 181, 0.22) !important;
            transition: all 0.25s cubic-bezier(0.16, 1, 0.3, 1) !important;
        }
        .stButton > button:hover, .stDownloadButton > button:hover {
            transform: translateY(-2px) !important;
            box-shadow: 0 8px 24px rgba(0, 240, 181, 0.4) !important;
            filter: brightness(1.08) !important;
        }
        .stButton > button:active, .stDownloadButton > button:active {
            transform: translateY(0) scale(.985) !important;
        }
        .stButton > button:focus-visible, .stDownloadButton > button:focus-visible {
            outline: 2px solid var(--geo-cyan) !important;
            outline-offset: 3px !important;
        }
        .stButton > button:disabled, .stDownloadButton > button:disabled {
            border-color: var(--geo-border) !important;
            background: #2a3944 !important;
            box-shadow: none !important;
            color: #a1afb8 !important;
            cursor: not-allowed !important;
        }

        /* Inputs, Forms & Selects */
        [data-testid="stTextInput"] input,
        [data-testid="stTextArea"] textarea,
        [data-testid="stNumberInput"] input {
            border: 1px solid var(--geo-border) !important;
            border-radius: 12px !important;
            background: #111c26 !important;
            color: var(--geo-text-head) !important;
            font-family: inherit !important;
            transition: all 0.25s ease !important;
        }
        [data-testid="stTextInput"] input:focus,
        [data-testid="stTextArea"] textarea:focus,
        [data-testid="stNumberInput"] input:focus {
            border-color: var(--geo-emerald) !important;
            box-shadow: 0 0 0 3px rgba(0, 240, 181, 0.16) !important;
        }
        [data-testid="stTextInput"] input::placeholder,
        [data-testid="stTextArea"] textarea::placeholder {
            color: #8296a4 !important;
            opacity: 1;
        }
        [data-testid="stSelectbox"] [data-baseweb="select"] > div,
        [data-testid="stMultiSelect"] [data-baseweb="select"] > div {
            transition: border-color .2s ease, box-shadow .2s ease, background .2s ease;
        }
        [data-testid="stSelectbox"] [data-baseweb="select"] > div:focus-within,
        [data-testid="stMultiSelect"] [data-baseweb="select"] > div:focus-within {
            border-color: var(--geo-emerald) !important;
            box-shadow: 0 0 0 3px rgba(0,240,181,.14) !important;
        }

        [data-testid="stSelectbox"] [data-baseweb="select"] > div,
        [data-testid="stMultiSelect"] [data-baseweb="select"] > div {
            border-color: var(--geo-border) !important;
            border-radius: 12px !important;
            background: #111c26 !important;
            color: var(--geo-text-head) !important;
        }

        [data-testid="stHorizontalBlock"] {
            gap: 1.25rem !important;
            margin-bottom: 0.85rem !important;
        }
        [data-testid="stVerticalBlockBorderWrapper"] {
            border: 1px solid var(--geo-border) !important;
            border-radius: 20px !important;
            background: var(--geo-card) !important;
            box-shadow: var(--geo-shadow-card) !important;
            backdrop-filter: blur(12px) !important;
        }
        [data-testid="stDataFrame"], [data-testid="stTable"] {
            border: 1px solid var(--geo-border) !important;
            border-radius: 16px !important;
            overflow: hidden !important;
            background: #111c26 !important;
            box-shadow: var(--geo-shadow-card) !important;
        }
        [data-testid="stCode"] {
            border: 1px solid var(--geo-border) !important;
            border-radius: 14px !important;
            background: #101b25 !important;
            font-family: ui-monospace, "Cascadia Code", Consolas, monospace !important;
        }
        [data-testid="stForm"], [data-testid="stFileUploader"] section,
        [data-testid="stExpander"] {
            border: 1px solid var(--geo-border) !important;
            border-radius: 16px !important;
            background: var(--geo-card) !important;
            backdrop-filter: blur(14px) !important;
        }
        [data-testid="stExpander"] details > summary {
            min-height: 2.8rem;
            padding-inline: .4rem;
            color: var(--geo-text-head);
            font-weight: 650;
        }
        [data-testid="stExpander"] summary:hover {
            color: var(--geo-emerald) !important;
        }
        [data-testid="stFileUploader"] {
            border: 1px dashed rgba(0, 240, 181, 0.35) !important;
            border-radius: 16px !important;
            background: rgba(0, 240, 181, 0.03) !important;
            transition: all 0.25s ease !important;
        }
        [data-testid="stFileUploader"]:hover {
            border-color: var(--geo-emerald) !important;
            background: rgba(0, 240, 181, 0.06) !important;
        }
        [data-baseweb="popover"], [data-baseweb="menu"] {
            background: #111c26 !important;
            border: 1px solid var(--geo-border) !important;
            border-radius: 14px !important;
            box-shadow: 0 16px 36px rgba(0, 0, 0, 0.5) !important;
        }
        [data-baseweb="menu"] li:hover {
            background: rgba(0, 240, 181, 0.12) !important;
            color: #FFFFFF !important;
        }
        [data-testid="stAlert"] {
            border: 1px solid var(--geo-border) !important;
            border-radius: 14px !important;
            background: rgba(20, 32, 43, 0.94) !important;
            box-shadow: var(--geo-shadow-card) !important;
        }
        [data-testid="stAlert"] p { color: var(--geo-text-body); }
        [data-testid="stAlert"]:has([data-testid="stAlertContentSuccess"]) {
            border-left: 3px solid var(--geo-emerald) !important;
        }
        [data-testid="stAlert"]:has([data-testid="stAlertContentWarning"]) {
            border-left: 3px solid var(--geo-amber) !important;
        }
        [data-testid="stAlert"]:has([data-testid="stAlertContentError"]) {
            border-left: 3px solid var(--geo-crimson) !important;
        }
        .saas-empty {
            padding: 2rem 1.5rem;
            border: 1px dashed rgba(189,207,220,.22);
            border-radius: var(--geo-radius-lg);
            background: linear-gradient(145deg, rgba(20,32,43,.86), rgba(17,28,38,.64));
            color: var(--geo-text-body);
            text-align: center;
        }
        .saas-empty-mark {
            display: inline-grid; place-items: center; width: 46px; height: 46px;
            margin-bottom: .7rem; border: 1px solid rgba(0,240,181,.2);
            border-radius: 15px; background: rgba(0,240,181,.08);
            color: var(--geo-emerald); font-size: 1.2rem;
        }
        .saas-empty-title { color: var(--geo-text-head); font-weight: 700; }
        .saas-empty-copy { margin-top: .35rem; color: var(--geo-text-muted); font-size: .83rem; }
        [data-testid="stSkeleton"] {
            border-radius: var(--geo-radius-md) !important;
            background: linear-gradient(100deg, #172631 25%, #243b48 42%, #172631 62%) !important;
            background-size: 220% 100% !important;
            animation: skeleton-shimmer 1.8s ease-in-out infinite !important;
        }
        @keyframes skeleton-shimmer {
            to { background-position-x: -220%; }
        }
        .stProgress > div > div > div > div {
            background: linear-gradient(90deg, var(--geo-emerald), var(--geo-cyan));
        }
        .stSpinner > div > div {
            border-top-color: var(--geo-emerald) !important;
        }

        /* 3D Slow-Motion Gyroscopic & Polyhedron Engine */
        .mine-hero-mark {
            position: relative;
            flex: 0 0 240px;
            width: 240px;
            height: 220px;
            perspective: 1100px;
        }
        .mine-3d-scene {
            position: relative;
            width: 100%;
            height: 100%;
            transform-style: preserve-3d;
            animation: scene-float 38s ease-in-out infinite alternate;
            will-change: transform;
        }
        .mine-orbit {
            position: absolute;
            left: 50%;
            top: 50%;
            border-radius: 50%;
            transform-style: preserve-3d;
            pointer-events: none;
        }
        .orbit-outer {
            width: 210px;
            height: 105px;
            margin-left: -105px;
            margin-top: -52px;
            border: 1px solid rgba(0, 240, 181, 0.38);
            box-shadow: 0 0 16px rgba(0, 240, 181, 0.15), inset 0 0 12px rgba(0, 240, 181, 0.1);
            animation: orbit-spin-1 56s linear infinite;
        }
        .orbit-mid {
            width: 170px;
            height: 85px;
            margin-left: -85px;
            margin-top: -42px;
            border: 1px solid rgba(0, 216, 246, 0.35);
            box-shadow: 0 0 14px rgba(0, 216, 246, 0.12);
            animation: orbit-spin-2 68s linear infinite reverse;
        }
        .orbit-inner {
            width: 130px;
            height: 65px;
            margin-left: -65px;
            margin-top: -32px;
            border: 1px dashed rgba(245, 186, 107, 0.45);
            animation: orbit-spin-3 82s linear infinite;
        }
        .orbit-sat {
            position: absolute;
            top: -5px;
            left: 50%;
            width: 9px;
            height: 9px;
            margin-left: -4.5px;
            border-radius: 50%;
        }
        .sat-emerald {
            background: #00F0B5;
            box-shadow: 0 0 10px #00F0B5, 0 0 22px rgba(0,240,181,0.8);
        }
        .sat-cyan {
            background: #00D8F6;
            box-shadow: 0 0 10px #00D8F6, 0 0 20px rgba(0,216,246,0.8);
        }
        .sat-copper {
            background: #F5BA6B;
            box-shadow: 0 0 8px #F5BA6B, 0 0 16px rgba(245,186,107,0.7);
        }
        .mine-polyhedron {
            position: absolute;
            left: 50%;
            top: 50%;
            width: 86px;
            height: 106px;
            margin-left: -43px;
            margin-top: -53px;
            transform-style: preserve-3d;
            animation: polyhedron-tumble 72s linear infinite;
        }
        .facet {
            position: absolute;
            inset: 0;
            border-radius: 14px;
            backdrop-filter: blur(5px);
        }
        .facet-1 {
            clip-path: polygon(50% 0%, 100% 30%, 82% 100%, 18% 100%, 0% 30%);
            background: linear-gradient(145deg, rgba(0, 240, 181, 0.48), rgba(0, 216, 246, 0.28) 50%, rgba(245, 186, 107, 0.45));
            border: 1px solid rgba(255, 255, 255, 0.35);
            box-shadow: inset 0 0 20px rgba(0, 240, 181, 0.35), 0 0 32px rgba(0, 240, 181, 0.25);
            transform: translateZ(20px) rotateY(16deg);
        }
        .facet-2 {
            clip-path: polygon(50% 0%, 100% 30%, 82% 100%, 18% 100%, 0% 30%);
            background: linear-gradient(135deg, rgba(245, 186, 107, 0.42), rgba(0, 240, 181, 0.24) 60%, rgba(0, 216, 246, 0.38));
            border: 1px solid rgba(255, 255, 255, 0.25);
            transform: translateZ(-20px) rotateY(-16deg) rotateX(180deg);
        }
        .facet-core {
            position: absolute;
            left: 50%;
            top: 50%;
            width: 26px;
            height: 26px;
            margin-left: -13px;
            margin-top: -13px;
            border-radius: 50%;
            background: radial-gradient(circle, #FFFFFF 0%, #00F0B5 55%, #00D8F6 100%);
            box-shadow: 0 0 24px #00F0B5, 0 0 48px rgba(0,240,181,0.65);
            animation: core-pulse 12s ease-in-out infinite alternate;
        }
        .mine-sparkle {
            position: absolute;
            color: #00F0B5;
            font-size: 13px;
            pointer-events: none;
            opacity: 0.85;
            filter: drop-shadow(0 0 6px #00F0B5);
        }
        .sparkle-1 { top: 12%; left: 8%; animation: sparkle-float-1 17s ease-in-out infinite; }
        .sparkle-2 { bottom: 18%; right: 10%; animation: sparkle-float-2 21s ease-in-out infinite 2s; }
        .sparkle-3 { top: 20%; right: 18%; animation: sparkle-float-1 24s ease-in-out infinite 4s; }

        @keyframes ambient-drift {
            0% { background-position: 0% 0%, 100% 100%, 0% 0%; }
            100% { background-position: 100% 70%, 0% 20%, 0% 0%; }
        }
        @keyframes hero-aura-drift {
            0% { background-position: 0% 50%, 100% 50%, 0% 0%; }
            100% { background-position: 100% 40%, 0% 60%, 0% 0%; }
        }
        @keyframes scene-float {
            0% { transform: translateY(0px) rotateX(2deg) rotateY(-3deg); }
            50% { transform: translateY(-9px) rotateX(-2deg) rotateY(4deg); }
            100% { transform: translateY(2px) rotateX(3deg) rotateY(-2deg); }
        }
        @keyframes orbit-spin-1 {
            0% { transform: rotateX(68deg) rotateZ(0deg); }
            100% { transform: rotateX(68deg) rotateZ(360deg); }
        }
        @keyframes orbit-spin-2 {
            0% { transform: rotateX(55deg) rotateY(25deg) rotateZ(0deg); }
            100% { transform: rotateX(55deg) rotateY(25deg) rotateZ(360deg); }
        }
        @keyframes orbit-spin-3 {
            0% { transform: rotateX(72deg) rotateY(-30deg) rotateZ(0deg); }
            100% { transform: rotateX(72deg) rotateY(-30deg) rotateZ(360deg); }
        }
        @keyframes polyhedron-tumble {
            0% { transform: rotateX(0deg) rotateY(0deg) rotateZ(0deg); }
            50% { transform: rotateX(180deg) rotateY(180deg) rotateZ(90deg); }
            100% { transform: rotateX(360deg) rotateY(360deg) rotateZ(360deg); }
        }
        @keyframes core-pulse {
            0% { transform: scale(0.82); opacity: 0.65; }
            100% { transform: scale(1.22); opacity: 1; box-shadow: 0 0 34px #00F0B5, 0 0 65px rgba(0,216,246,0.85); }
        }
        @keyframes sparkle-float-1 {
            0%, 100% { transform: translate(0, 0) scale(0.7); opacity: 0.25; }
            50% { transform: translate(8px, -14px) scale(1.25); opacity: 1; }
        }
        @keyframes sparkle-float-2 {
            0%, 100% { transform: translate(0, 0) scale(0.6); opacity: 0.2; }
            50% { transform: translate(-10px, -16px) scale(1.15); opacity: 0.95; }
        }
        @keyframes pulse-soft {
            0%, 100% { box-shadow: 0 0 0 4px rgba(0,240,181,.18); }
            50% { box-shadow: 0 0 0 8px rgba(0,240,181,.06); }
        }

        @media (max-width: 760px) {
            .mine-hero { min-height: 230px; padding: 1.45rem; }
            .mine-hero-mark { flex-basis: 115px; width: 115px; transform: scale(.7); margin-right: -1.1rem; }
            .mine-workflow { grid-template-columns: repeat(2, minmax(0, 1fr)); }
        }
        @media (max-width: 480px) {
            .mine-hero { align-items: flex-start; }
            .mine-hero-mark { position: absolute; right: -1.6rem; bottom: -1rem; opacity: .48; }
            .mine-hero-copy { position: relative; z-index: 2; }
            .mine-workflow { grid-template-columns: 1fr 1fr; }
            [data-testid="stMainBlockContainer"] {
                padding: 1rem .8rem 2rem !important;
            }
            .saas-kpi { min-height: 132px; padding: 1rem; }
        }
        @media (prefers-reduced-motion: reduce) {
            *, *::before, *::after {
                animation-duration: .01ms !important;
                animation-iteration-count: 1 !important;
                scroll-behavior: auto !important;
                transition-duration: .01ms !important;
            }
        }
        </style>
        """,
        unsafe_allow_html=True,
    )

    workflow_labels = [
        ("01 · INGEST", "Upload & extract"),
        ("02 · NORMALIZE", "Domain facts"),
        ("03 · VALIDATE", "Consistency firewall"),
        ("04 · TRACE", "Evidence search"),
        ("05 · APPROVE", "Human decision"),
    ]
    workflow = "".join(
        f'<div class="mine-step{" is-active" if index == active_step else ""}">'
        f"<span>{label}</span>{description}</div>"
        for index, (label, description) in enumerate(workflow_labels)
    )
    st.markdown(
        f"""
        <header class="mine-hero">
          <div class="mine-hero-copy">
            <div class="mine-eyebrow"><i class="mine-live-dot"></i> {icon} &nbsp; {active_page} · GeoMine intelligence</div>
            <h1>{heading}</h1>
            <p>{description}</p>
            <div class="mine-hero-foot">
              <span class="mine-chip">Fact passports</span>
              <span class="mine-chip">Deterministic checks</span>
              <span class="mine-chip">Human-approved output</span>
            </div>
          </div>
          <div class="mine-hero-mark" aria-hidden="true">
            <div class="mine-3d-scene">
              <div class="mine-orbit orbit-outer">
                <div class="orbit-sat sat-emerald"></div>
              </div>
              <div class="mine-orbit orbit-mid">
                <div class="orbit-sat sat-cyan"></div>
              </div>
              <div class="mine-orbit orbit-inner">
                <div class="orbit-sat sat-copper"></div>
              </div>
              <div class="mine-polyhedron">
                <div class="facet facet-1"></div>
                <div class="facet facet-2"></div>
                <div class="facet-core"></div>
              </div>
              <div class="mine-sparkle sparkle-1">✦</div>
              <div class="mine-sparkle sparkle-2">✦</div>
              <div class="mine-sparkle sparkle-3">✧</div>
            </div>
          </div>
        </header>
        <div class="mine-workflow" aria-label="Reporting workflow">
          {workflow}
        </div>
        """,
        unsafe_allow_html=True,
    )


def render_kpi_card(label: str, value: int, icon: str, detail: str, state: str = "default") -> None:
    st.markdown(
        f"""
        <article class="saas-kpi saas-kpi-{state}">
          <div class="saas-kpi-top">
            <span class="saas-kpi-label">{label}</span>
            <span class="saas-kpi-icon" aria-hidden="true">{icon}</span>
          </div>
          <div class="saas-kpi-value">{value:,}</div>
          <div class="saas-kpi-detail">{detail}</div>
          <div class="saas-kpi-accent"></div>
        </article>
        """,
        unsafe_allow_html=True,
    )


def render_empty_state(title: str, message: str, icon: str = "◇") -> None:
    st.markdown(
        f"""
        <div class="saas-empty">
          <div class="saas-empty-mark">{icon}</div>
          <div class="saas-empty-title">{title}</div>
          <div class="saas-empty-copy">{message}</div>
        </div>
        """,
        unsafe_allow_html=True,
    )


def main():
    st.set_page_config(page_title="CMPDI/CIL Unique AI Reporting", page_icon="⛏️", layout="wide")

    init_db()
    load_ontology_overrides()

    with st.sidebar:
        st.markdown(
            """
            <div class="sidebar-brand">
              <div class="sidebar-brand-mark">⛏</div>
              <div class="sidebar-brand-title">GeoMine Intelligence</div>
              <div class="sidebar-brand-caption">Evidence-led reporting workspace</div>
            </div>
            """,
            unsafe_allow_html=True,
        )
        page_icons = {
            "Dashboard": "◈",
            "Document Center": "▤",
            "Fact Passport": "◇",
            "Consistency Firewall": "⬡",
            "Change Intelligence": "↗",
            "AI Query Assistant": "⌕",
            "Parliamentary Copilot": "§",
            "Topic Intelligence": "◎",
            "Report Generator": "▧",
            "Mining Ontology": "⌘",
            "Audit Log": "≋",
        }
        page = st.radio(
            "WORKSPACE",
            options=[
                "Dashboard",
                "Document Center",
                "Fact Passport",
                "Consistency Firewall",
                "Change Intelligence",
                "AI Query Assistant",
                "Parliamentary Copilot",
                "Topic Intelligence",
                "Report Generator",
                "Mining Ontology",
                "Audit Log",
            ],
            format_func=lambda name: f"{page_icons[name]}   {name}",
        )

        st.divider()
        st.markdown(
            "**Core USP**  \n"
            "Every important fact can be traced from the final answer/report "
            "back to its source evidence."
        )

    render_visual_system(page)

    # ============================================================
    # DASHBOARD
    # ============================================================

    if page == "Dashboard":
        docs = documents_df()
        facts = facts_df()
        fw = consistency_firewall()
        conn = get_db()
        reports_generated = int(conn.execute("SELECT COUNT(*) FROM reports").fetchone()[0])

        verified_count = (
            int((facts["verification_status"] == "Verified").sum()) if not facts.empty else 0
        )
        first_row = st.columns(3, gap="medium")
        for column, metric in zip(
            first_row,
            [
                ("Documents indexed", len(docs), "▤", "Source files in this workspace", "default"),
                (
                    "Facts extracted",
                    len(facts),
                    "◇",
                    "Structured values with provenance",
                    "default",
                ),
                (
                    "Verified facts",
                    verified_count,
                    "✓",
                    "Reviewed by an authorized user",
                    "default",
                ),
            ],
            strict=True,
        ):
            with column:
                render_kpi_card(*metric)

        second_row = st.columns(3, gap="medium")
        for column, metric in zip(
            second_row,
            [
                (
                    "Hard conflicts",
                    fw["errors"],
                    "!",
                    "Must be resolved before final approval",
                    "danger" if fw["errors"] else "default",
                ),
                (
                    "Review warnings",
                    fw["warnings"],
                    "△",
                    "Low confidence or period checks",
                    "warning" if fw["warnings"] else "default",
                ),
                (
                    "Reports created",
                    reports_generated,
                    "▧",
                    "Draft and approved report records",
                    "default",
                ),
            ],
            strict=True,
        ):
            with column:
                render_kpi_card(*metric)

        st.subheader("Workspace overview")
        overview, firewall_column = st.columns([1.45, 1], gap="medium")
        with overview:
            with st.container(border=True):
                st.markdown("#### Reporting workflow")
                st.caption(
                    "Each stage builds on traceable evidence; numerical answers "
                    "come from structured facts, not generated guesses."
                )
                st.code(
                    "UPLOAD  →  EXTRACT  →  NORMALIZE  →  VALIDATE  →  TRACE  →  APPROVE",
                    language="text",
                )
                st.markdown(
                    "**Fact Passport**  ·  **Change Intelligence**  ·  **Evidence-backed reports**"
                )
        with firewall_column:
            with st.container(border=True):
                st.markdown("#### Consistency status")
                if fw["status"] == "PASS":
                    st.success("PASS · No hard conflicts detected")
                elif fw["status"] == "NO_DATA":
                    st.info("No indexed facts to validate yet")
                else:
                    st.error(f"REVIEW REQUIRED · {fw['errors']} hard conflict(s)")
                st.caption(
                    f"{fw['warnings']} warning(s) need review. "
                    "Warnings alone do not replace authorized approval."
                )

        with st.container(border=True):
            st.markdown("#### Intelligence modules")
            innovation = pd.DataFrame(
                [
                    [
                        "Fact Passport",
                        "Every figure retains source, page, snippet, confidence, and status.",
                    ],
                    [
                        "Consistency Firewall",
                        "Surfaces cross-document conflicts and arithmetic anomalies.",
                    ],
                    [
                        "Change Intelligence",
                        "Compares reporting periods and flags high-impact changes.",
                    ],
                    [
                        "Parliamentary Copilot",
                        "Creates sourced drafts that require human approval.",
                    ],
                    [
                        "Mining Ontology",
                        "Maps domain terminology to standard metrics and units.",
                    ],
                    [
                        "Evidence Trace",
                        "Links each answer to its fact and originating source.",
                    ],
                ],
                columns=["Module", "What it does"],
            )
            st.dataframe(innovation, use_container_width=True, hide_index=True)

        recent = pd.read_sql_query(
            "SELECT event_type,message,created_at FROM audits ORDER BY id DESC LIMIT 8",
            conn,
        )
        conn.close()
        recent_facts = (
            facts.head(8)[
                [
                    "fact_id",
                    "entity",
                    "metric",
                    "period",
                    "value",
                    "unit",
                    "verification_status",
                ]
            ]
            if not facts.empty
            else pd.DataFrame()
        )
        activity_column, facts_column = st.columns([1, 1.4], gap="medium")
        with activity_column:
            with st.container(border=True):
                st.markdown("#### Recent activity")
                if recent.empty:
                    st.caption("Activity will appear here as records are processed.")
                else:
                    st.dataframe(recent, use_container_width=True, hide_index=True)
        with facts_column:
            with st.container(border=True):
                st.markdown("#### Latest fact passports")
                if recent_facts.empty:
                    st.caption("Index a source document to create the first fact passport.")
                else:
                    st.dataframe(recent_facts, use_container_width=True, hide_index=True)

    # ============================================================
    # DOCUMENT CENTER
    # ============================================================

    elif page == "Document Center":
        st.header("📄 Document Center")
        st.write(
            "Upload approved demo/non-confidential documents. "
            "The app extracts domain facts and stores source provenance."
        )

        uploads = st.file_uploader(
            "Files",
            type=["pdf", "docx", "xlsx", "csv", "txt", "png", "jpg", "jpeg"],
            accept_multiple_files=True,
        )

        if uploads and st.button("Process & Index Documents", type="primary"):
            total = 0
            with st.spinner("Extracting documents and building evidence passports..."):
                for upload in uploads:
                    try:
                        inserted, fact_count = ingest_document(upload.getvalue(), upload.name)
                        if inserted:
                            st.success(f"{upload.name}: indexed, {fact_count} facts extracted.")
                            total += fact_count
                        else:
                            st.info(f"{upload.name}: already indexed.")
                    except (OSError, RuntimeError, ValueError, pd.errors.ParserError) as exc:
                        st.error(f"{upload.name}: {exc}")

            st.info(f"Batch extraction complete. Facts detected: {total}")

        if st.button("Index synthetic demo dataset"):
            demo_path = APP_DIR / "demo_data" / "synthetic_mine_report.csv"
            try:
                inserted, fact_count = ingest_document(demo_path.read_bytes(), demo_path.name)
                if inserted:
                    audit("DEMO_DATA_INGESTED", f"{demo_path.name}: {fact_count} facts")
                    st.success(f"Synthetic demo data indexed: {fact_count} facts.")
                else:
                    st.info("Synthetic demo data is already indexed.")
            except (OSError, RuntimeError, ValueError, pd.errors.ParserError) as exc:
                st.error(f"Could not index synthetic demo data: {exc}")

        docs = documents_df()
        if docs.empty:
            render_empty_state(
                "Your document library is ready",
                "Upload source files above or load the synthetic demo dataset to begin.",
                "▤",
            )
        else:
            st.subheader("Indexed documents")
            display_docs = docs[["id", "filename", "file_type", "uploaded_at", "metadata"]].copy()
            display_docs["table_count"] = display_docs["metadata"].apply(
                lambda raw: json.loads(raw or "{}").get("table_count", 0)
            )
            st.dataframe(
                display_docs.drop(columns=["metadata"]), use_container_width=True, hide_index=True
            )

    # ============================================================
    # FACT PASSPORT
    # ============================================================

    elif page == "Fact Passport":
        st.header("🪪 Fact Passport")
        df = facts_df()

        if df.empty:
            render_empty_state(
                "No fact passports yet",
                "Index a source document to see extracted values, provenance, and verification.",
                "◇",
            )
        else:
            fid = st.selectbox("Fact", df["fact_id"].tolist())
            row = df[df["fact_id"] == fid].iloc[0]

            c1, c2, c3 = st.columns(3)
            c1.metric(
                "Value", f"{float(row['value']):,.4f}".rstrip("0").rstrip(".") + f" {row['unit']}"
            )
            c2.metric("Confidence", f"{float(row['confidence']) * 100:.1f}%")
            c3.metric("Verification", row["verification_status"])

            st.json(build_fact_passport(row))

            if row["verification_status"] != "Verified":
                approver_name = st.text_input("Authorized approver name")
                approver_pin = st.text_input("Approver PIN", type="password")
                pin_is_configured = bool(os.environ.get("CMPDI_APPROVER_PIN"))
                if not pin_is_configured:
                    st.warning(
                        "Fact verification is disabled until CMPDI_APPROVER_PIN is configured."
                    )
                if st.button(
                    "Mark Fact as Verified",
                    type="primary",
                    disabled=not pin_is_configured,
                ):
                    if not verify_fact(fid, approver_name, approver_pin):
                        st.error("Approver credentials were not accepted.")
                    else:
                        st.success("Fact verified.")
                        st.rerun()
            else:
                st.success("This fact has been verified by an authorized user.")

    # ============================================================
    # CONSISTENCY FIREWALL
    # ============================================================

    elif page == "Consistency Firewall":
        st.header("🛡️ Report Consistency Firewall")
        record_validation = st.button("Run validation and record audit event")
        result = consistency_firewall(record_audit=record_validation)

        if result["status"] == "PASS":
            st.success("PASS — no hard source conflicts detected.")
        elif result["status"] == "REVIEW_REQUIRED":
            st.error("REVIEW REQUIRED — source conflicts detected.")
        elif result["status"] == "NO_DATA":
            st.info("NO DATA — index source documents before validation.")

        c1, c2 = st.columns(2)
        c1.metric("Errors", result["errors"])
        c2.metric("Warnings", result["warnings"])

        for item in result["findings"]:
            if item["severity"] == "ERROR":
                st.error(f"{item['type']}: {item['message']}\n\n{item['detail']}")
            else:
                st.warning(f"{item['type']}: {item['message']}\n\n{item['detail']}")

    # ============================================================
    # CHANGE INTELLIGENCE
    # ============================================================

    elif page == "Change Intelligence":
        st.header("📈 Change Intelligence")
        df = facts_df()

        if df.empty:
            render_empty_state(
                "No period data to compare",
                "Index records covering at least two reporting periods to unlock change intelligence.",
                "↗",
            )
        else:
            entities = sorted(df["entity"].dropna().unique().tolist())
            entity = st.selectbox("Entity", entities)

            periods = sorted(df[df["entity"] == entity]["period"].dropna().unique().tolist())

            if len(periods) < 2:
                render_empty_state(
                    "One more reporting period needed",
                    "This entity needs facts from at least two distinct reporting periods.",
                    "◷",
                )
            else:
                old_period = st.selectbox("Older period", periods, index=0)
                new_period = st.selectbox("Newer period", periods, index=1)

                if st.button("Run Change Intelligence", type="primary"):
                    out = compare_entity_periods(entity, old_period, new_period)
                    st.dataframe(out, use_container_width=True)

                    attention = out[out["Attention"] == "ATTENTION"]
                    if not attention.empty:
                        st.warning(
                            f"{len(attention)} high-impact change(s) detected. "
                            "Review source evidence before interpretation."
                        )
                        st.dataframe(attention, use_container_width=True)

    # ============================================================
    # AI QUERY ASSISTANT
    # ============================================================

    elif page == "AI Query Assistant":
        st.header("💬 AI Query Assistant")
        st.write(
            "This prototype answers from indexed evidence. Structured facts are preferred "
            "for numerical answers."
        )

        query = st.text_input(
            "Question", placeholder="What was production of Mine Alpha in FY2024-25?"
        )

        use_ollama = st.checkbox("Optional: use local Ollama for answer wording")
        model = st.text_input("Ollama model", "llama3.2") if use_ollama else ""

        if st.button("Ask", type="primary") and query.strip():
            result = evidence_answer(query)
            st.markdown(result["answer"])

            if result["fact"] is not None:
                st.subheader("Evidence Trace")
                for fact in result.get("facts", [result["fact"]]):
                    st.json(build_fact_passport(fact))

            st.subheader("Retrieved source documents")
            for src in result["sources"]:
                st.markdown(f"**{src['filename']}** — relevance {src.get('relevance', 0):.3f}")
                snippet = " ".join(str(src["content"]).split())
                st.caption(snippet[:800] + ("..." if len(snippet) > 800 else ""))

            if use_ollama:
                context = {
                    "fact": build_fact_passport(result["fact"])
                    if result["fact"] is not None
                    else None,
                    "sources": [
                        {"filename": x["filename"], "content": str(x["content"])[:5000]}
                        for x in result["sources"]
                    ],
                }
                prompt = (
                    "Rewrite the answer using only the supplied evidence. "
                    "Do not invent numbers, dates, entities or sources.\n\n"
                    f"Question: {query}\nEvidence:\n{json.dumps(context, default=str)}"
                )
                try:
                    wording = ollama_rewrite(prompt, model)
                    st.subheader("Local AI wording")
                    st.write(wording)
                    st.caption(
                        "Wording only: numerical values and source references remain "
                        "the deterministic Fact Passport result above."
                    )
                except (RuntimeError, ValueError) as exc:
                    st.error(str(exc))

    # ============================================================
    # PARLIAMENTARY COPILOT
    # ============================================================

    elif page == "Parliamentary Copilot":
        st.header("🏛️ Parliamentary Question Copilot")

        question = st.text_area(
            "Paste the question",
            placeholder=(
                "State the production of Mine Alpha for the last three financial years "
                "and percentage change."
            ),
            height=150,
        )

        if st.button("Analyse & Draft", type="primary") and question.strip():
            analysis, draft, result = draft_parliamentary_answer(question)

            st.subheader("Question decomposition")
            st.json(analysis)

            st.subheader("Draft")
            st.markdown(draft)

            st.subheader("Evidence")
            if result["fact"] is not None:
                for fact in result.get("facts", [result["fact"]]):
                    st.json(build_fact_passport(fact))
            else:
                st.warning("No structured fact matched the question.")

            fw = consistency_firewall()
            if fw["status"] == "PASS":
                st.success("Current knowledge base passes the hard-conflict check.")
            else:
                st.warning(
                    f"Current knowledge base contains {fw['errors']} conflict(s). "
                    "Review before official use."
                )

    # ============================================================
    # TOPIC INTELLIGENCE
    # ============================================================

    elif page == "Topic Intelligence":
        st.header("☁️ Topic Intelligence & Word Cloud")
        docs = documents_df()

        if docs.empty:
            render_empty_state(
                "Topic intelligence is waiting for records",
                "Index mining and geology documents to discover keywords and evolving themes.",
                "◎",
            )
        else:
            selected = st.multiselect(
                "Documents", docs["filename"].tolist(), default=docs["filename"].tolist()[:5]
            )
            text = "\n".join(
                docs.loc[docs["filename"].isin(selected), "content"].fillna("").tolist()
            )

            scores = domain_topic_scores(text)
            st.subheader("Domain topic ranking")
            st.dataframe(
                pd.DataFrame(
                    [(k, v) for k, v in scores.items() if v > 0], columns=["Topic", "Keyword Score"]
                ),
                use_container_width=True,
                hide_index=True,
            )

            st.subheader("Top keywords")
            st.dataframe(
                pd.DataFrame(top_keywords(text), columns=["Keyword", "Frequency"]),
                use_container_width=True,
                hide_index=True,
            )

            if WordCloud is not None and plt is not None and text.strip():
                try:
                    wc = WordCloud(
                        width=1200,
                        height=500,
                        background_color="white",
                        stopwords=STOPWORDS,
                    ).generate(text)
                    fig, ax = plt.subplots(figsize=(14, 5))
                    ax.imshow(wc, interpolation="bilinear")
                    ax.axis("off")
                    st.pyplot(fig, clear_figure=True)
                except ValueError as exc:
                    st.warning(f"Word cloud is unavailable for the selected text: {exc}")

            period_facts = facts_df()
            if not period_facts.empty and selected:
                selected_doc_ids = docs.loc[docs["filename"].isin(selected), "id"].tolist()
                selected_facts = period_facts[period_facts["document_id"].isin(selected_doc_ids)]
                period_rows = []
                for period, group in selected_facts.groupby("period"):
                    doc_content = docs.loc[docs["id"].isin(group["document_id"]), "content"].fillna(
                        ""
                    )
                    topic_scores = domain_topic_scores("\n".join(doc_content))
                    for topic, score in topic_scores.items():
                        if score:
                            period_rows.append(
                                {"Period": period, "Topic": topic, "Frequency": score}
                            )
                if period_rows:
                    st.subheader("Topic evolution by reporting period")
                    st.dataframe(
                        pd.DataFrame(period_rows),
                        use_container_width=True,
                        hide_index=True,
                    )

    # ============================================================
    # REPORT GENERATOR
    # ============================================================

    elif page == "Report Generator":
        st.header("📝 Automated Report Generator")
        df = facts_df()

        if df.empty:
            render_empty_state(
                "Nothing to report yet",
                "Extract and validate facts before creating a traceable report.",
                "▧",
            )
        else:
            title = st.text_input(
                "Report title", value="CMPDI/CIL AI-Assisted Geological & Mining Report"
            )

            entities = sorted(df["entity"].dropna().unique().tolist())
            selected_entities = st.multiselect("Entities", entities, default=entities[:3])

            selected = (
                df[df["entity"].isin(selected_entities)] if selected_entities else df.head(100)
            )

            compare_entity = st.selectbox(
                "Optional comparison entity",
                ["None"] + entities,
                key="report_compare_entity",
            )
            comparison = pd.DataFrame()
            if compare_entity != "None":
                compare_periods = sorted(
                    df.loc[df["entity"] == compare_entity, "period"].dropna().unique()
                )
                if len(compare_periods) >= 2:
                    c_old, c_new = st.columns(2)
                    old_period = c_old.selectbox(
                        "Comparison: older period", compare_periods, key="report_old_period"
                    )
                    new_choices = [period for period in compare_periods if period != old_period]
                    new_period = c_new.selectbox(
                        "Comparison: newer period", new_choices, key="report_new_period"
                    )
                    comparison = compare_entity_periods(compare_entity, old_period, new_period)
                    st.subheader("Period comparison")
                    st.dataframe(comparison, use_container_width=True, hide_index=True)
                else:
                    st.info("The selected entity needs at least two periods for comparison.")

            st.subheader("Selected evidence")
            st.dataframe(
                selected[
                    [
                        "fact_id",
                        "entity",
                        "metric",
                        "period",
                        "value",
                        "unit",
                        "confidence",
                        "filename",
                        "source_page",
                    ]
                ].head(50),
                use_container_width=True,
            )

            fw = consistency_firewall()
            st.info(
                f"Firewall: {fw['status']} | Errors: {fw['errors']} | Warnings: {fw['warnings']}"
            )

            c1, c2 = st.columns(2)

            with c1:
                if st.button("Generate DOCX", type="primary"):
                    try:
                        with st.spinner("Preparing the evidence-linked DOCX report..."):
                            file_bytes = generate_docx(title, selected, fw, comparison)
                            report_id = record_generated_report(
                                title, "DOCX", fw, selected, comparison
                            )
                        st.download_button(
                            "Download DOCX",
                            data=file_bytes,
                            file_name="CMPDI_CIL_AI_Report.docx",
                            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        )
                        st.success(
                            f"Draft report #{report_id} generated. Human approval is required."
                        )
                    except Exception as e:
                        st.error(str(e))

            with c2:
                if st.button("Generate PDF"):
                    try:
                        with st.spinner("Preparing the evidence-linked PDF report..."):
                            file_bytes = generate_pdf(title, selected, fw, comparison)
                            report_id = record_generated_report(
                                title, "PDF", fw, selected, comparison
                            )
                        st.download_button(
                            "Download PDF",
                            data=file_bytes,
                            file_name="CMPDI_CIL_AI_Report.pdf",
                            mime="application/pdf",
                        )
                        st.success(
                            f"Draft report #{report_id} generated. Human approval is required."
                        )
                    except Exception as e:
                        st.error(str(e))

            st.subheader("Human approval workflow")
            if fw["errors"]:
                st.error(
                    "Final approval is blocked while hard consistency conflicts exist. "
                    "Resolve the source conflicts and generate a new report."
                )
            elif not os.environ.get("CMPDI_APPROVER_PIN"):
                st.warning(
                    "Set CMPDI_APPROVER_PIN in the deployment environment to enable "
                    "authorized report approval."
                )
            conn = get_db()
            report_history = pd.read_sql_query(
                """SELECT id,title,file_type,status,generated_at,approved_at,approved_by,
                          firewall_status
                   FROM reports ORDER BY id DESC""",
                conn,
            )
            conn.close()
            if not report_history.empty:
                st.dataframe(report_history, use_container_width=True, hide_index=True)
                approvable = report_history[
                    (report_history["status"] == "Draft")
                    & (report_history["firewall_status"] == "PASS")
                ]
                if not approvable.empty:
                    selected_report_id = st.selectbox(
                        "Draft to approve",
                        approvable["id"].astype(int).tolist(),
                    )
                    approver_name = st.text_input("Approver name", key="report_approver_name")
                    approver_pin = st.text_input(
                        "Approver PIN", type="password", key="report_approver_pin"
                    )
                    if st.button(
                        "Approve final report",
                        disabled=not bool(os.environ.get("CMPDI_APPROVER_PIN")),
                    ):
                        if approve_report(selected_report_id, approver_name, approver_pin):
                            st.success("Report approved and approval event audited.")
                            st.rerun()
                        else:
                            st.error(
                                "Approval failed. Verify credentials and confirm the report "
                                "passed the Consistency Firewall."
                            )
            approved_reports = report_history[report_history["status"] == "Approved"]
            if not approved_reports.empty:
                st.subheader("Final approved outputs")
                approved_id = st.selectbox(
                    "Approved report",
                    approved_reports["id"].astype(int).tolist(),
                    key="approved_report_output",
                )
                conn = get_db()
                approved_row = conn.execute(
                    "SELECT * FROM reports WHERE id=?", (approved_id,)
                ).fetchone()
                conn.close()
                saved_ids = json.loads(approved_row["fact_ids"])
                approved_facts = facts[facts["fact_id"].isin(saved_ids)]
                saved_comparison = pd.DataFrame(json.loads(approved_row["comparison_data"]))
                saved_firewall = json.loads(approved_row["firewall_data"])
                approval_note = (
                    f"APPROVED by {approved_row['approved_by']} on {approved_row['approved_at']}."
                )
                final_docx, final_pdf = st.columns(2)
                with final_docx:
                    if st.button("Generate approved DOCX"):
                        output = generate_docx(
                            approved_row["title"],
                            approved_facts,
                            saved_firewall,
                            saved_comparison,
                            approval_note,
                        )
                        audit("FINAL_REPORT_GENERATED", f"DOCX report #{approved_id}")
                        st.download_button(
                            "Download approved DOCX",
                            data=output,
                            file_name=f"Approved_Report_{approved_id}.docx",
                            mime=(
                                "application/vnd.openxmlformats-officedocument."
                                "wordprocessingml.document"
                            ),
                        )
                with final_pdf:
                    if st.button("Generate approved PDF"):
                        output = generate_pdf(
                            approved_row["title"],
                            approved_facts,
                            saved_firewall,
                            saved_comparison,
                            approval_note,
                        )
                        audit("FINAL_REPORT_GENERATED", f"PDF report #{approved_id}")
                        st.download_button(
                            "Download approved PDF",
                            data=output,
                            file_name=f"Approved_Report_{approved_id}.pdf",
                            mime="application/pdf",
                        )
    # ============================================================
    # MINING ONTOLOGY
    # ============================================================

    elif page == "Mining Ontology":
        st.header("⛏️ Mining / Geological Domain Ontology")

        st.write(
            "The purpose of the ontology is to make generic document AI domain-aware. "
            "It links technical terms to standard metrics and query concepts."
        )

        rows = []
        for key, meta in ONTOLOGY.items():
            rows.append(
                {
                    "Metric": key,
                    "Label": meta["label"],
                    "Aliases": ", ".join(meta["aliases"]),
                    "Expected units": ", ".join(meta["units"]),
                }
            )
        st.dataframe(pd.DataFrame(rows), use_container_width=True, hide_index=True)

        term = st.text_input("Test a mining/geological term", placeholder="e.g. GCV")
        if term:
            matches = []
            t = term.lower().strip()
            for key, meta in ONTOLOGY.items():
                if t == key or t in [a.lower() for a in meta["aliases"]]:
                    matches.append(
                        {"Metric": key, "Meaning": meta["label"], "Units": meta["units"]}
                    )

            if matches:
                st.success("Ontology match found.")
                st.json(matches)
            else:
                st.warning("No exact ontology match. Add it to the domain dictionary.")

        st.subheader("Add or update a terminology mapping")
        with st.form("ontology_entry_form"):
            metric_key = st.text_input("Metric key", placeholder="e.g. seam_thickness")
            metric_label = st.text_input("Display label", placeholder="Seam Thickness")
            aliases = st.text_input("Aliases, comma-separated")
            units = st.text_input("Standard units, comma-separated")
            save_entry = st.form_submit_button("Save ontology mapping")
        if save_entry:
            key = re.sub(r"[^a-z0-9_]+", "_", metric_key.lower()).strip("_")
            alias_values = [value.strip() for value in aliases.split(",") if value.strip()]
            unit_values = [value.strip() for value in units.split(",") if value.strip()]
            if not key or not metric_label.strip() or not alias_values or not unit_values:
                st.error(
                    "Metric key, display label, at least one alias, and one unit are required."
                )
            else:
                ONTOLOGY[key] = {
                    "label": metric_label.strip(),
                    "aliases": alias_values,
                    "units": unit_values,
                }
                try:
                    save_ontology()
                    audit("ONTOLOGY_UPDATED", f"{key}: {', '.join(alias_values)}")
                    st.success(f"Saved {key} to {ONTOLOGY_PATH.name}.")
                    st.rerun()
                except OSError as exc:
                    st.error(f"Could not save ontology configuration: {exc}")

    # ============================================================
    # AUDIT
    # ============================================================

    elif page == "Audit Log":
        st.header("🔎 Audit Log")
        conn = get_db()
        logs = pd.read_sql_query(
            "SELECT event_type,message,created_at FROM audits ORDER BY id DESC", conn
        )
        conn.close()

        if logs.empty:
            st.info("No audit events yet.")
        else:
            st.dataframe(logs, use_container_width=True, hide_index=True)

    st.divider()
    st.caption(
        "Prototype only. Use approved/synthetic data. AI outputs are drafts and "
        "must be validated through the authorized CMPDI/CIL workflow."
    )


if __name__ == "__main__":
    main()
